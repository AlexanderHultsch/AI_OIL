"""Teams connector: chats and channel posts captured by the Power Automate bridge (CON-60 .. CON-68)."""

import fnmatch
import io
import json
import os
import re
from collections.abc import Mapping
from dataclasses import dataclass, replace
from datetime import UTC, datetime
from typing import ClassVar
from urllib.parse import unquote, urlsplit

import openpyxl
from openpyxl.utils import range_boundaries

from ai_oil.connectors import email_text
from ai_oil.connectors.sharepoint import _CONFLICT_SUFFIX, FileSystem, LocalFileSystem
from ai_oil.errors import ValidationError
from ai_oil.sources.model import (
    CandidateQuery,
    ConnectorConfig,
    EntryData,
    FetchResult,
    HealthCheck,
    HealthReport,
    IdentityRef,
    ObjectMeta,
    ObjectState,
    SourceRef,
)
from ai_oil.timeutil import local_display, parse_iso

_REQUIRED_COLUMNS = ("MessageId", "ConversationId", "ConversationType", "Timestamp", "SenderName", "Text")
_OFFLINE_SKIPPED = "not available offline — skipped (hydrate is off)"
_HYDRATED = "hydrated: downloaded on read (Files On-Demand)"
_TEAMS_HOSTS = ("teams.microsoft.com", "teams.cloud.microsoft")
_RAW_PREFIX = "teams:"
_MAX_NAMED_SENDERS = 4
_EARLIEST = datetime.min.replace(tzinfo=UTC)
_LOCK_FILE_PREFIX = "~$"

_NOT_A_MESSAGE_LINK = (
    "only Teams message links can be linked — in Teams, open the menu (…) of a message in the chat or"
    " channel and choose Copy link"
)


@dataclass(frozen=True)
class ConversationCandidate:
    """A captured conversation found by search or scan (CON-67)."""

    ref: SourceRef
    title: str
    latest: datetime
    sender: str
    message_count: int


@dataclass(frozen=True)
class _Row:
    message_id: str
    conversation_id: str
    conversation_type: str
    conversation_name: str
    sent_at: datetime
    sender_name: str
    sender_email: str
    permalink: str
    text: str


@dataclass(frozen=True)
class _FileData:
    name: str
    rows: tuple[_Row, ...]
    complete: bool
    problem: str | None
    notes: tuple[str, ...]

    @property
    def problems(self) -> tuple[str, ...]:
        return (f"{self.name}: {self.problem}",) if self.problem else ()

    @property
    def notices(self) -> tuple[str, ...]:
        return tuple(f"{self.name}: {note}" for note in self.notes)


@dataclass(frozen=True)
class _Bridge:
    files: tuple[_FileData, ...]
    conflicts: tuple[str, ...]
    conversations: Mapping[str, list[_Row]]
    messages: Mapping[str, _Row]

    @property
    def complete(self) -> bool:
        return all(file.complete for file in self.files)

    @property
    def problems(self) -> tuple[str, ...]:
        """Files that could not be used; reported on every fetch."""
        return tuple(problem for file in self.files for problem in file.problems)

    @property
    def notices(self) -> tuple[str, ...]:
        """Conflict copies, offline files, skipped rows and naive timestamps; reported once (CON-66)."""
        conflicts = tuple(f"conflict copy of a bridge file: {name} (not read)" for name in self.conflicts)
        return (*conflicts, *(notice for file in self.files for notice in file.notices))

    def newest(self) -> datetime | None:
        return max((rows[-1].sent_at for rows in self.conversations.values()), default=None)


def _iso(moment: datetime) -> str:
    utc = moment.astimezone(UTC)
    return f"{utc:%Y-%m-%dT%H:%M:%S}.{utc.microsecond // 1000:03d}Z"


def _setting_text(settings: Mapping[str, object], key: str, default: str | None) -> str:
    value = settings.get(key, default)
    if value is None:
        raise ValidationError(f"Teams setting {key} is required")
    if not isinstance(value, str) or not value.strip():
        raise ValidationError(f"Teams setting {key} must be non-blank text")
    return value.strip()


def _bool_setting(settings: Mapping[str, object], key: str, default: bool) -> bool:
    value = settings.get(key, default)
    if not isinstance(value, bool):
        raise ValidationError(f"Teams setting {key} must be true or false")
    return value


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def _timestamp(value: object) -> tuple[datetime | None, bool]:
    """The moment in UTC (None if unreadable) and whether the cell carried no time zone (CON-62)."""
    moment: datetime | None = None
    if isinstance(value, datetime):
        moment = value
    elif isinstance(value, str):
        try:
            moment = datetime.fromisoformat(value.strip())
        except ValueError:
            return None, False
    if moment is None:
        return None, False
    if moment.tzinfo is None:
        return moment.replace(tzinfo=UTC), True
    return moment.astimezone(UTC), False


def _order(row: _Row) -> tuple[datetime, str]:
    return row.sent_at, row.message_id


def _merge(first: _Row, other: _Row) -> _Row:
    """The same message written by both flows: the first row wins, gaps are filled from the other (CON-63)."""
    return replace(
        first,
        conversation_name=first.conversation_name or other.conversation_name,
        sender_email=first.sender_email or other.sender_email,
        permalink=first.permalink or other.permalink,
    )


def _title(rows: list[_Row]) -> str:
    """The newest conversation name, else the senders of a chat, else a fixed channel title (CON-64)."""
    named = next((row.conversation_name for row in reversed(rows) if row.conversation_name), None)
    if named:
        return named
    if rows[-1].conversation_type.casefold() == "channel":
        return "Channel conversation"
    senders = list(dict.fromkeys(row.sender_name for row in rows if row.sender_name))
    if not senders:
        return "Chat"
    shown = ", ".join(senders[:_MAX_NAMED_SENDERS])
    more = len(senders) - _MAX_NAMED_SENDERS
    return f"Chat with {shown} and {more} more" if more > 0 else f"Chat with {shown}"


def _locator(rows: list[_Row]) -> str:
    link = next((row.permalink for row in reversed(rows) if row.permalink), None)
    return link or f"{_RAW_PREFIX}{rows[0].conversation_id}"


class TeamsConnector:
    """Teams conversations read from the Excel files of the Power Automate bridge (CON-6, CON-60)."""

    type: ClassVar[str] = "teams"
    object_types: ClassVar[frozenset[str]] = frozenset({"teams_thread"})
    action_class: ClassVar[str] = "A"
    max_parallel: ClassVar[int] = 1

    def __init__(self, config: ConnectorConfig, *, fs: FileSystem | None = None) -> None:
        self.config = config
        settings = config.settings
        self._folder = _setting_text(settings, "folder", None)
        self._pattern = _setting_text(settings, "file_pattern", "Teams_Bridge*.xlsx")
        self._table = _setting_text(settings, "table", "TeamsMessages")
        self._hydrate = _bool_setting(settings, "hydrate", True)
        self._fs: FileSystem = fs if fs is not None else LocalFileSystem()
        self._files: dict[tuple[str, float, int], _FileData] = {}
        self._plain: dict[str, str] = {}
        self._reported: set[str] = set()

    # --- reading the bridge ----------------------------------------------------------------------------

    def _bridge_names(self) -> tuple[list[str], list[str]]:
        """The bridge files to read and the conflict copies among them, sorted by name (CON-61, CON-66)."""
        wanted = self._pattern.casefold()
        names = sorted(
            name
            for name in self._fs.list_dir(self._folder)
            if not name.startswith(_LOCK_FILE_PREFIX)
            and fnmatch.fnmatchcase(name.casefold(), wanted)
            and self._fs.is_file(os.path.join(self._folder, name))
        )
        stems = [os.path.splitext(name) for name in names]
        conflicts = []
        for name, (stem, extension) in zip(names, stems, strict=True):
            if any(
                other != name
                and other_extension.lower() == extension.lower()
                and re.fullmatch(re.escape(other_stem) + _CONFLICT_SUFFIX, stem)
                for other, (other_stem, other_extension) in zip(names, stems, strict=True)
            ):
                conflicts.append(name)
        return [name for name in names if name not in conflicts], conflicts

    def _read_file(self, name: str) -> _FileData:
        path = os.path.join(self._folder, name)
        info = self._fs.info(path)
        key = (path, info.mtime, info.size)
        if key in self._files:
            return self._files[key]
        notes = []
        if info.offline:
            if not self._hydrate:
                return _FileData(name, (), False, None, (_OFFLINE_SKIPPED,))
            notes.append(_HYDRATED)
        parsed = self._parse(name, self._fs.read_bytes(path), notes)
        self._files[key] = parsed
        return parsed

    def _parse(self, name: str, data: bytes, notes: list[str]) -> _FileData:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception as exc:
            return _FileData(name, (), False, f"cannot read the workbook: {exc}", tuple(notes))
        try:
            return self._parse_table(name, workbook, notes)
        finally:
            workbook.close()

    def _parse_table(self, name: str, workbook: openpyxl.Workbook, notes: list[str]) -> _FileData:
        wanted = self._table.casefold()
        found = next(
            (
                (sheet, table)
                for sheet in workbook.worksheets
                for table in sheet.tables.values()
                if table.name.casefold() == wanted
            ),
            None,
        )
        if found is None:
            return _FileData(name, (), False, f"table {self._table} not found", tuple(notes))
        sheet, table = found
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        lines = sheet.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True
        )
        header = next(lines, ())
        columns: dict[str, int] = {}
        for index, cell in enumerate(header):
            columns.setdefault(_cell_text(cell).casefold(), index)
        missing = [column for column in _REQUIRED_COLUMNS if column.casefold() not in columns]
        if missing:
            problem = f"table {self._table} lacks the required column(s): {', '.join(missing)}"
            return _FileData(name, (), False, problem, tuple(notes))

        def cell(line: tuple, column: str) -> object:
            index = columns.get(column.casefold())
            return line[index] if index is not None and index < len(line) else None

        rows = []
        skipped = naive = 0
        for line in lines:
            message_id = _cell_text(cell(line, "MessageId"))
            conversation_id = _cell_text(cell(line, "ConversationId"))
            sent_at, was_naive = _timestamp(cell(line, "Timestamp"))
            if not message_id or not conversation_id or sent_at is None:
                skipped += 1
                continue
            naive += was_naive
            rows.append(
                _Row(
                    message_id,
                    conversation_id,
                    _cell_text(cell(line, "ConversationType")),
                    _cell_text(cell(line, "ConversationName")),
                    sent_at,
                    _cell_text(cell(line, "SenderName")),
                    _cell_text(cell(line, "SenderEmail")),
                    _cell_text(cell(line, "Permalink")),
                    _cell_text(cell(line, "Text")),
                )
            )
        if skipped:
            notes.append(f"{skipped} row(s) skipped: no MessageId, no ConversationId or no valid Timestamp")
        if naive:
            notes.append("timestamps without a time zone were taken as UTC")
        return _FileData(name, tuple(rows), True, None, tuple(notes))

    def _bridge(self) -> _Bridge:
        if not self._fs.is_dir(self._folder):
            raise ValidationError(
                f"Teams bridge folder not found: {self._folder} — is OneDrive running and the folder synced?"
            )
        names, conflicts = self._bridge_names()
        files = tuple(self._read_file(name) for name in names)
        messages: dict[str, _Row] = {}
        for file in files:
            for row in file.rows:
                known = messages.get(row.message_id)
                messages[row.message_id] = row if known is None else _merge(known, row)
        conversations: dict[str, list[_Row]] = {}
        for row in messages.values():
            conversations.setdefault(row.conversation_id, []).append(row)
        for rows in conversations.values():
            rows.sort(key=_order)
        return _Bridge(files, tuple(conflicts), conversations, messages)

    def check(self) -> HealthReport:
        if not self._fs.is_dir(self._folder):
            return HealthReport((HealthCheck("folder", False, f"folder not found: {self._folder}"),))
        bridge = self._bridge()
        self._unreported(bridge.notices)
        count = len(bridge.files)
        checks = [
            HealthCheck("folder", True, self._folder),
            HealthCheck(
                "files",
                count > 0,
                f"{count} file(s) matching {self._pattern}"
                if count
                else f"no file matching {self._pattern} in {self._folder}",
            ),
        ]
        for file in bridge.files:
            if file.complete:
                detail = f"table {self._table}: {len(file.rows)} message(s)"
                detail += "".join(f"; {note}" for note in file.notes)
            else:
                detail = "; ".join(part for part in (file.problem, *file.notes) if part)
            checks.append(HealthCheck(f"file {file.name}", file.complete, detail))
        checks += [
            HealthCheck(f"conflict copy {name}", False, "conflict copy of a bridge file (not read)")
            for name in bridge.conflicts
        ]
        newest = bridge.newest()
        checks.append(
            HealthCheck(
                "newest message",
                True,
                local_display(_iso(newest)) if newest is not None else "none captured yet",
            )
        )
        return HealthReport(tuple(checks))

    # --- references ------------------------------------------------------------------------------------

    def _ref(self, conversation_id: str, locator: str, rows: list[_Row] | None) -> SourceRef:
        title = _title(rows) if rows else None
        return SourceRef("teams_thread", conversation_id, display_ref=title, locator=locator)

    def parse_reference(self, text: str) -> SourceRef | None:
        candidate = text.strip()
        if candidate.casefold().startswith(_RAW_PREFIX):
            conversation_id = candidate[len(_RAW_PREFIX) :].strip()
            if not conversation_id:
                return None
            return self._ref(conversation_id, candidate, self._bridge().conversations.get(conversation_id))
        parts = urlsplit(candidate)
        if parts.scheme != "https" or (parts.hostname or "").lower() not in _TEAMS_HOSTS:
            return None
        segments = [unquote(segment) for segment in parts.path.split("/") if segment]
        if len(segments) < 4 or [segment.lower() for segment in segments[:2]] != ["l", "message"]:
            raise ValidationError(_NOT_A_MESSAGE_LINK)
        link_conversation, message_id = segments[2], segments[3]
        bridge = self._bridge()
        row = bridge.messages.get(message_id)
        conversation_id = row.conversation_id if row is not None else link_conversation
        return self._ref(conversation_id, candidate, bridge.conversations.get(conversation_id))

    def resolve(self, ref: SourceRef) -> ObjectMeta:
        if ref.object_type != "teams_thread":
            raise ValidationError(f"the Teams connector cannot resolve a {ref.object_type} reference")
        return ObjectMeta(
            "teams_thread",
            ref.external_id,
            title=ref.display_ref,
            display_ref=ref.display_ref,
            locator=ref.locator,
        )

    # --- searching and scanning ------------------------------------------------------------------------

    def _plain_text(self, html: str) -> str:
        if html not in self._plain:
            self._plain[html] = email_text.html_to_text(html)
        return self._plain[html]

    def _candidate(self, rows: list[_Row]) -> ConversationCandidate:
        newest = rows[-1]
        title = _title(rows)
        return ConversationCandidate(
            ref=SourceRef("teams_thread", newest.conversation_id, display_ref=title, locator=_locator(rows)),
            title=title,
            latest=newest.sent_at,
            sender=newest.sender_name or newest.sender_email or "-",
            message_count=len(rows),
        )

    def _matches(self, rows: list[_Row], needle: str) -> bool:
        return (
            needle in _title(rows).casefold()
            or any(needle in row.sender_name.casefold() for row in rows)
            or any(needle in self._plain_text(row.text).casefold() for row in rows)
        )

    def find_conversations(self, text: str, *, limit: int = 20) -> list[ConversationCandidate]:
        """Conversations whose name, senders or messages contain `text`, newest first (CON-67)."""
        needle = text.strip().casefold()
        found = [
            self._candidate(rows)
            for rows in self._bridge().conversations.values()
            if self._matches(rows, needle)
        ]
        return sorted(
            found, key=lambda candidate: (candidate.latest, candidate.ref.external_id), reverse=True
        )[:limit]

    def recent_conversations(self, *, since: datetime) -> list[ConversationCandidate]:
        """Conversations with messages at or after `since`, newest first, for the scan (CON-67)."""
        if since.tzinfo is None:
            since = since.replace(tzinfo=UTC)
        found = [
            self._candidate(rows)
            for rows in self._bridge().conversations.values()
            if rows[-1].sent_at >= since
        ]
        return sorted(
            found, key=lambda candidate: (candidate.latest, candidate.ref.external_id), reverse=True
        )

    def search_candidates(self, query: CandidateQuery) -> list[ObjectMeta]:
        since = parse_iso(query.since) if query.since else None
        return [
            ObjectMeta(
                "teams_thread",
                candidate.ref.external_id,
                title=candidate.title,
                display_ref=candidate.title,
                locator=candidate.ref.locator,
            )
            for candidate in self.find_conversations(query.text, limit=query.limit)
            if since is None or candidate.latest >= since
        ]

    # --- fetching --------------------------------------------------------------------------------------

    def _entry(self, row: _Row) -> EntryData:
        email = row.sender_email.lower()
        author = IdentityRef("email", email, row.sender_name or None) if "@" in email else None
        metadata: dict[str, object] = {"sender_name": row.sender_name}
        if row.sender_email and author is None:
            metadata["sender_raw"] = row.sender_email
        if row.permalink:
            metadata["permalink"] = row.permalink
        return EntryData(
            "teams_message",
            row.message_id,
            _iso(row.sent_at),
            self._plain_text(row.text),
            author=author,
            metadata=metadata,
        )

    def _unreported(self, notices: tuple[str, ...]) -> list[str]:
        """The notices not yet reported by this instance, i.e. in this sync run (CON-66)."""
        fresh = [notice for notice in dict.fromkeys(notices) if notice not in self._reported]
        self._reported.update(fresh)
        return fresh

    def fetch(self, state: ObjectState) -> FetchResult:
        bridge = self._bridge()
        rows = bridge.conversations.get(state.external_id)
        warnings = [*bridge.problems, *self._unreported(bridge.notices)]
        if not rows:
            if not bridge.complete:
                warnings.append(
                    "conversation not found, but not every bridge file could be read — not marked as gone"
                )
                return FetchResult(cursor=state.cursor, warnings=tuple(warnings))
            warnings.append("conversation not found in the bridge files (not captured yet, or removed)")
            return FetchResult(status="gone", warnings=tuple(warnings))

        return FetchResult(
            entries=tuple(self._entry(row) for row in rows),
            cursor=json.dumps({"latest": _iso(rows[-1].sent_at)}),
            title=_title(rows),
            warnings=tuple(warnings),
        )
