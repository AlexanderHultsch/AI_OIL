"""Text extraction for versioning and search (spec 6.7)."""

import email
import email.policy
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath

import docx
import extract_msg
import openpyxl
import pptx
import pypdf
from bs4 import BeautifulSoup

SUPPORTED_EXTENSIONS = (
    ".docx",
    ".xlsx",
    ".xlsm",
    ".pptx",
    ".pdf",
    ".msg",
    ".eml",
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
)

_PLAIN_TEXT_EXTENSIONS = (".txt", ".md", ".csv", ".json", ".xml")
_NO_TEXT_DETAIL = "no extractable text (scanned PDF?); OCR is not performed"


@dataclass(frozen=True)
class ExtractResult:
    text: str
    status: str  # ok | no_text | unsupported | error
    truncated: bool = False
    detail: str | None = None


def extract_text(
    data: bytes, filename: str, *, max_chars: int = 2_000_000, max_cells: int = 50_000
) -> ExtractResult:
    """Turn file bytes into plain text; never raises for bad input."""
    suffix = PurePath(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return ExtractResult(
            "", "unsupported", detail=f"no extractor for {suffix or 'files without extension'}"
        )
    try:
        text, truncated = _extract(data, suffix, max_cells)
    except Exception as exc:
        return ExtractResult("", "error", detail=f"{type(exc).__name__}: {exc}")
    if suffix == ".pdf" and not text.strip():
        return ExtractResult("", "no_text", detail=_NO_TEXT_DETAIL)
    if len(text) > max_chars:
        text = text[:max_chars]
        truncated = True
    return ExtractResult(text, "ok", truncated=truncated)


def _extract(data: bytes, suffix: str, max_cells: int) -> tuple[str, bool]:
    if suffix == ".docx":
        return _docx_text(data), False
    if suffix in (".xlsx", ".xlsm"):
        return _xlsx_text(data, max_cells)
    if suffix == ".pptx":
        return _pptx_text(data), False
    if suffix == ".pdf":
        return _pdf_text(data), False
    if suffix == ".msg":
        return _msg_text(data), False
    if suffix == ".eml":
        return _eml_text(data), False
    return data.decode("utf-8-sig", errors="replace"), False


def _docx_text(data: bytes) -> str:
    document = docx.Document(BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _xlsx_text(data: bytes, max_cells: int) -> tuple[str, bool]:
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    cells = 0
    truncated = False
    try:
        for sheet in workbook.worksheets:
            lines.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                while values and values[-1] == "":
                    values.pop()
                if not any(values):
                    continue
                if cells >= max_cells:
                    truncated = True
                    break
                if len(values) > max_cells - cells:
                    values = values[: max_cells - cells]
                    truncated = True
                cells += len(values)
                lines.append("\t".join(values))
            if truncated:
                break
    finally:
        workbook.close()
    return "\n".join(lines), truncated


def _pptx_text(data: bytes) -> str:
    presentation = pptx.Presentation(BytesIO(data))
    lines: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines.append(f"## Slide {number}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
    return "\n".join(lines)


def _pdf_text(data: bytes) -> str:
    reader = pypdf.PdfReader(BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _msg_text(data: bytes) -> str:
    message = extract_msg.openMsg(data)
    try:
        body = message.body
        if not (body or "").strip() and message.htmlBody:
            body = _html_to_text(message.htmlBody)
        return _mail_text(message.sender, message.to, message.cc, message.date, message.subject, body)
    finally:
        message.close()


def _eml_text(data: bytes) -> str:
    message = email.message_from_bytes(data, policy=email.policy.default)
    part = message.get_body(preferencelist=("plain", "html"))
    body = ""
    if part is not None:
        body = part.get_content()
        if part.get_content_type() == "text/html":
            body = _html_to_text(body)
    return _mail_text(
        message["From"], message["To"], message["Cc"], message["Date"], message["Subject"], body
    )


def _html_to_text(html: str | bytes) -> str:
    return BeautifulSoup(html, "html.parser").get_text("\n")


def _mail_text(sender: object, to: object, cc: object, date: object, subject: object, body: object) -> str:
    lines = [f"From: {sender or ''}", f"To: {to or ''}"]
    if cc:
        lines.append(f"Cc: {cc}")
    lines.append(f"Date: {date or ''}")
    lines.append(f"Subject: {subject or ''}")
    lines.append("")
    lines.append(str(body or ""))
    return "\n".join(lines)
