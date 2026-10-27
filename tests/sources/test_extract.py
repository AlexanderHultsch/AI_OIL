import datetime
from email.message import EmailMessage
from io import BytesIO

import docx
import extract_msg
import openpyxl
import pptx
import pytest

from ai_oil.sources.extract import SUPPORTED_EXTENSIONS, extract_text


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    table = document.add_table(rows=2, cols=2)
    for row_index, values in enumerate((("a", "b"), ("c", "d"))):
        for col_index, value in enumerate(values):
            table.cell(row_index, col_index).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Sheet1"
    first.append(["x", None, "z", None])
    first.append([1, 2.5, 3])
    second = workbook.create_sheet("Parts")
    second.append(["Inductor", "lead time"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    presentation = pptx.Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Gate review"
    first.placeholders[1].text = "Open points"
    first.notes_slide.notes_text_frame.text = "Mention the PPAP delay"
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Next steps"
    second.placeholders[1].text = "Freeze BOM"
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(content: bytes) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref_at)
    return out


def _eml_bytes(*, plain: str | None = None, html: str | None = None) -> bytes:
    message = EmailMessage()
    message["From"] = "alice@example.com"
    message["To"] = "bob@example.com"
    message["Cc"] = "carol@example.com"
    message["Date"] = "Mon, 26 Oct 2026 09:00:00 +0000"
    message["Subject"] = "Inductor lead time"
    if plain is not None:
        message.set_content(plain)
        if html is not None:
            message.add_alternative(html, subtype="html")
    else:
        message.set_content(html or "", subtype="html")
    return message.as_bytes()


class _FakeMsg:
    def __init__(
        self, body: str | None, html_body: bytes | None, cc: str | None = "carol@example.com"
    ) -> None:
        self.sender = "alice@example.com"
        self.to = "bob@example.com"
        self.cc = cc
        self.date = datetime.datetime(2026, 10, 26, 9, 0, tzinfo=datetime.UTC)
        self.subject = "Inductor lead time"
        self.body = body
        self.htmlBody = html_body
        self.closed = False

    def close(self) -> None:
        self.closed = True


def test_supported_extensions_cover_the_spec_formats() -> None:
    assert ".docx" in SUPPORTED_EXTENSIONS
    assert ".msg" in SUPPORTED_EXTENSIONS
    assert ".zip" not in SUPPORTED_EXTENSIONS


def test_docx_paragraphs_then_table_rows() -> None:
    result = extract_text(_docx_bytes(), "spec.DOCX")
    assert result.status == "ok"
    assert not result.truncated
    assert "First paragraph\nSecond paragraph" in result.text
    assert "a\tb" in result.text
    assert "c\td" in result.text
    assert result.text.index("Second paragraph") < result.text.index("a\tb")


def test_xlsx_sheets_rows_and_trailing_empties() -> None:
    result = extract_text(_xlsx_bytes(), "plan.xlsx")
    assert result.status == "ok"
    assert result.text.splitlines() == [
        "## Sheet1",
        "x\t\tz",
        "1\t2.5\t3",
        "## Parts",
        "Inductor\tlead time",
    ]


def test_xlsx_stops_after_max_cells_and_flags_truncation() -> None:
    result = extract_text(_xlsx_bytes(), "plan.xlsm", max_cells=4)
    assert result.status == "ok"
    assert result.truncated
    assert "Inductor" not in result.text
    assert result.text.splitlines()[0] == "## Sheet1"


def test_pptx_slides_texts_and_notes() -> None:
    result = extract_text(_pptx_bytes(), "deck.pptx")
    assert result.status == "ok"
    lines = result.text.splitlines()
    assert lines[0] == "## Slide 1"
    assert "Gate review" in lines
    assert "Open points" in lines
    assert "Notes: Mention the PPAP delay" in lines
    assert "## Slide 2" in lines
    assert "Freeze BOM" in lines
    assert result.text.count("Notes:") == 1


def test_pdf_text_is_extracted() -> None:
    content = b"BT /F1 18 Tf 20 100 Td (Inductor lead time) Tj ET"
    result = extract_text(_pdf_bytes(content), "supplier.pdf")
    assert result.status == "ok"
    assert "Inductor lead time" in result.text


def test_pdf_without_text_is_flagged() -> None:
    result = extract_text(_pdf_bytes(b""), "scan.pdf")
    assert result.status == "no_text"
    assert result.text == ""
    assert result.detail == "no extractable text (scanned PDF?); OCR is not performed"


def test_eml_plain_body_with_header_block() -> None:
    result = extract_text(_eml_bytes(plain="Please confirm 12 weeks."), "mail.eml")
    assert result.status == "ok"
    assert result.text.splitlines()[:6] == [
        "From: alice@example.com",
        "To: bob@example.com",
        "Cc: carol@example.com",
        "Date: Mon, 26 Oct 2026 09:00:00 +0000",
        "Subject: Inductor lead time",
        "",
    ]
    assert "Please confirm 12 weeks." in result.text


def test_eml_html_only_body_has_tags_removed() -> None:
    result = extract_text(
        _eml_bytes(html="<html><body><p>Confirm <b>12</b> weeks.</p></body></html>"), "m.eml"
    )
    assert result.status == "ok"
    assert "Subject: Inductor lead time" in result.text
    assert "Confirm" in result.text
    assert "12" in result.text
    assert "<" not in result.text


def test_eml_prefers_plain_over_html() -> None:
    result = extract_text(_eml_bytes(plain="plain wins", html="<p>html loses</p>"), "m.eml")
    assert "plain wins" in result.text
    assert "html loses" not in result.text


def test_eml_without_cc_omits_the_line() -> None:
    message = EmailMessage()
    message["From"] = "a@example.com"
    message["To"] = "b@example.com"
    message["Subject"] = "Hi"
    message.set_content("body")
    result = extract_text(message.as_bytes(), "m.eml")
    assert "Cc:" not in result.text


def test_msg_plain_body(monkeypatch: pytest.MonkeyPatch) -> None:
    fake = _FakeMsg("Please confirm 12 weeks.", None)
    seen: list[bytes] = []

    def fake_open(data: bytes) -> _FakeMsg:
        seen.append(data)
        return fake

    monkeypatch.setattr(extract_msg, "openMsg", fake_open)
    result = extract_text(b"raw-msg-bytes", "mail.msg")
    assert seen == [b"raw-msg-bytes"]
    assert fake.closed
    assert result.status == "ok"
    assert result.text.splitlines() == [
        "From: alice@example.com",
        "To: bob@example.com",
        "Cc: carol@example.com",
        "Date: 2026-10-26 09:00:00+00:00",
        "Subject: Inductor lead time",
        "",
        "Please confirm 12 weeks.",
    ]


def test_msg_html_only_body(monkeypatch: pytest.MonkeyPatch) -> None:
    fake = _FakeMsg(None, b"<html><body><p>Confirm <b>12</b> weeks.</p></body></html>", cc=None)
    monkeypatch.setattr(extract_msg, "openMsg", lambda data: fake)
    result = extract_text(b"x", "mail.msg")
    assert fake.closed
    assert result.status == "ok"
    assert "Cc:" not in result.text
    assert "Confirm" in result.text
    assert "<" not in result.text


def test_msg_is_closed_when_reading_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    class Broken(_FakeMsg):
        @property
        def sender(self) -> str:
            raise ValueError("boom")

        @sender.setter
        def sender(self, value: str) -> None:
            pass

    fake = Broken("body", None)
    monkeypatch.setattr(extract_msg, "openMsg", lambda data: fake)
    result = extract_text(b"x", "mail.msg")
    assert fake.closed
    assert result.status == "error"
    assert result.detail == "ValueError: boom"


def test_txt_with_bom_and_invalid_byte_is_decoded_with_replacement() -> None:
    result = extract_text(b"\xef\xbb\xbfcaf\xc3\xa9 \xff end", "notes.txt")
    assert result.status == "ok"
    assert result.text == "café � end"


def test_json_passes_through() -> None:
    payload = '{"part": "L1", "weeks": 12}'
    result = extract_text(payload.encode(), "data.json")
    assert result == extract_text(payload.encode(), "data.JSON")
    assert result.text == payload
    assert result.status == "ok"


def test_empty_supported_file_is_still_ok() -> None:
    result = extract_text(b"", "empty.md")
    assert result.status == "ok"
    assert result.text == ""


def test_unsupported_extension() -> None:
    result = extract_text(b"PK", "archive.zip")
    assert result.status == "unsupported"
    assert result.text == ""
    assert result.detail == "no extractor for .zip"


def test_file_without_extension_is_unsupported() -> None:
    result = extract_text(b"x", "README")
    assert result.status == "unsupported"
    assert result.detail == "no extractor for files without extension"


def test_corrupt_docx_is_an_error_not_an_exception() -> None:
    result = extract_text(b"this is not a zip file", "broken.docx")
    assert result.status == "error"
    assert result.text == ""
    assert result.detail


def test_max_chars_truncates_text() -> None:
    result = extract_text(b"abcdefghij", "notes.txt", max_chars=4)
    assert result.status == "ok"
    assert result.text == "abcd"
    assert result.truncated
    assert not extract_text(b"abcd", "notes.txt", max_chars=4).truncated
