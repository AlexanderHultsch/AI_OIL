import os
import re
from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
import sqlalchemy

from ai_oil.config import load_config, write_default_config
from ai_oil.connectors.sharepoint import FileInfo, LocalFileSystem, SharePointConnector
from ai_oil.domain import issues, people, projects
from ai_oil.errors import ValidationError
from ai_oil.sources import linking, registry, store
from ai_oil.sources.model import (
    CandidateQuery,
    ConnectorConfig,
    FetchResult,
    ObjectMeta,
    ObjectState,
    SourceRef,
)
from ai_oil.sources.sync import SyncScope, run_sync

LIBRARY_URL = "https://acme.sharepoint.com/sites/HW/Shared Documents"
ENCODED_LIBRARY_URL = "https://acme.sharepoint.com/sites/HW/Shared%20Documents"
MTIME = 1_700_000_000
OFFLINE_SKIPPED = "not available offline — skipped (hydrate is off)"
HYDRATED = "hydrated: downloaded on fetch (Files On-Demand)"


class FakeFileSystem:
    """The local file system with a marker for Files On-Demand placeholders and a log of the reads."""

    def __init__(self) -> None:
        self._local = LocalFileSystem()
        self.offline: set[str] = set()
        self.reads: list[str] = []

    def is_file(self, path: str) -> bool:
        return self._local.is_file(path)

    def is_dir(self, path: str) -> bool:
        return self._local.is_dir(path)

    def info(self, path: str) -> FileInfo:
        real = self._local.info(path)
        return FileInfo(real.mtime, real.size, path in self.offline)

    def read_bytes(self, path: str) -> bytes:
        self.reads.append(path)
        return self._local.read_bytes(path)

    def list_dir(self, path: str) -> list[str]:
        return self._local.list_dir(path)

    def walk_files(self, root: str):
        return self._local.walk_files(root)


def config(settings: dict) -> ConnectorConfig:
    return ConnectorConfig(1, "sharepoint", "SharePoint", True, settings, None)


def roots_setting(*roots: tuple[str, Path]) -> dict:
    return {"sync_roots": [{"url": url, "path": str(path)} for url, path in roots]}


@pytest.fixture
def library(tmp_path: Path) -> Path:
    path = tmp_path / "OneDrive" / "HW - Shared Documents"
    (path / "Specs").mkdir(parents=True)
    return path


@pytest.fixture
def fs() -> FakeFileSystem:
    return FakeFileSystem()


def make(library: Path, fs: FakeFileSystem | None = None, **settings: object) -> SharePointConnector:
    return SharePointConnector(config({**roots_setting((LIBRARY_URL, library)), **settings}), fs=fs)


def write(path: Path, text: str = "Lead time is 26 weeks.", mtime: int = MTIME) -> str:
    path.write_text(text, encoding="utf-8")
    os.utime(path, (mtime, mtime))
    return str(path)


def state(path: str, cursor: str | None = None) -> ObjectState:
    return ObjectState(
        id=1,
        connector_id=1,
        object_type="file",
        external_id=path,
        title=None,
        display_ref=os.path.basename(path),
        locator=path,
        cursor=cursor,
        parent_id=None,
        latest_content_hash=None,
    )


def fetch(connector: SharePointConnector, path: str, cursor: str | None = None) -> FetchResult:
    return connector.fetch(state(path, cursor))


# --- settings and contract -------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("settings", "key"),
    [
        ({"sync_roots": "C:\\Sync"}, "sync_roots"),
        ({"sync_roots": ["C:\\Sync"]}, "sync_roots"),
        ({"sync_roots": [{"url": LIBRARY_URL}]}, "sync_roots"),
        ({"sync_roots": [{"path": "/sync"}]}, "sync_roots"),
        ({"sync_roots": [{"url": " ", "path": "/sync"}]}, "sync_roots"),
        ({"sync_roots": [{"url": LIBRARY_URL, "path": ""}]}, "sync_roots"),
        ({"sync_roots": [{"url": 5, "path": "/sync"}]}, "sync_roots"),
        ({"hydrate": "yes"}, "hydrate"),
        ({"max_text_chars": "many"}, "max_text_chars"),
        ({"max_text_chars": -1}, "max_text_chars"),
        ({"max_text_chars": True}, "max_text_chars"),
    ],
)
def test_invalid_settings_are_rejected_naming_the_key(settings: dict, key: str) -> None:
    with pytest.raises(ValidationError, match=key):
        SharePointConnector(config(settings))


def test_the_defaults_need_no_settings() -> None:
    connector = SharePointConnector(config({}))

    assert isinstance(connector._fs, LocalFileSystem)
    assert connector._roots == []
    assert connector._hydrate is True
    assert connector._max_text_chars == 2_000_000


def test_the_connector_declares_its_contract() -> None:
    assert SharePointConnector.type == "sharepoint"
    assert SharePointConnector.object_types == frozenset({"file"})
    assert SharePointConnector.action_class == "A"
    assert SharePointConnector.max_parallel == 4


# --- the local file system -------------------------------------------------------------------------------


def test_the_local_file_system_reads_real_files(tmp_path: Path) -> None:
    (tmp_path / "sub").mkdir()
    one = write(tmp_path / "one.txt", "abc")
    two = write(tmp_path / "sub" / "two.txt", "abcdef", mtime=MTIME + 5)
    local = LocalFileSystem()

    assert (local.is_file(one), local.is_file(str(tmp_path / "sub")), local.is_file(one + "x")) == (
        True,
        False,
        False,
    )
    assert (local.is_dir(str(tmp_path)), local.is_dir(one)) == (True, False)
    assert local.info(two) == FileInfo(MTIME + 5, 6, False)
    assert local.read_bytes(one) == b"abc"
    assert sorted(local.list_dir(str(tmp_path))) == ["one.txt", "sub"]
    assert sorted(local.walk_files(str(tmp_path))) == sorted([one, two])
    assert list(local.walk_files(str(tmp_path / "missing"))) == []


@pytest.mark.parametrize(
    ("attributes", "offline"),
    [
        (0, False),
        (0x20, False),
        (0x1000, True),
        (0x40000, True),
        (0x400000, True),
        (0x20 | 0x400000, True),
    ],
)
def test_the_offline_flag_comes_from_the_windows_file_attributes(
    attributes: int, offline: bool, monkeypatch: pytest.MonkeyPatch
) -> None:
    stat = SimpleNamespace(st_mtime=12.5, st_size=7, st_file_attributes=attributes)
    monkeypatch.setattr(os, "stat", lambda path: stat)

    assert LocalFileSystem().info("any") == FileInfo(12.5, 7, offline)


def test_a_platform_without_file_attributes_is_never_offline(tmp_path: Path) -> None:
    path = write(tmp_path / "plain.txt")

    assert LocalFileSystem().info(path).offline is False


# --- parse_reference -------------------------------------------------------------------------------------


def test_a_local_path_to_a_file(library: Path) -> None:
    path = write(library / "Specs" / "torque.txt")

    ref = make(library).parse_reference(f"  {path}\n")

    assert ref == SourceRef("file", path, display_ref="torque.txt", locator=path)


def test_a_local_path_is_normalised(library: Path) -> None:
    path = write(library / "Specs" / "torque.txt")
    untidy = str(library / "Specs" / ".." / "Specs" / "torque.txt")

    ref = make(library).parse_reference(untidy)

    assert ref == SourceRef("file", path, display_ref="torque.txt", locator=path)


def test_an_absolute_path_that_is_not_a_file_is_not_a_reference(library: Path) -> None:
    assert make(library).parse_reference(str(library / "Specs" / "missing.txt")) is None
    assert make(library).parse_reference(str(library / "Specs")) is None


@pytest.mark.parametrize("text", ["torque.txt", "Specs/torque.txt", "", "OIL-1", "ftp://x.sharepoint.com/a"])
def test_other_text_is_not_a_reference(text: str, library: Path) -> None:
    assert make(library).parse_reference(text) is None


@pytest.mark.parametrize("base", [LIBRARY_URL, ENCODED_LIBRARY_URL])
def test_a_library_url_maps_to_the_synced_file(base: str, library: Path) -> None:
    path = str(library / "Specs" / "Torque spec.docx")
    url = f"{base}/Specs/Torque%20spec.docx"

    ref = make(library).parse_reference(url)

    assert ref == SourceRef("file", path, display_ref="Torque spec.docx", locator=url)


def test_the_file_does_not_have_to_exist_to_be_mapped(library: Path) -> None:
    url = f"{ENCODED_LIBRARY_URL}/Specs/later.docx"

    ref = make(library).parse_reference(url)

    assert ref is not None
    assert ref.external_id == str(library / "Specs" / "later.docx")


def test_a_sharing_prefix_with_a_query_string_is_removed(library: Path) -> None:
    url = (
        "https://acme.sharepoint.com/:x:/r/sites/HW/Shared%20Documents/Specs/Plan.xlsx"
        "?d=w1234&csf=1&web=1&e=AbCd#frag"
    )

    ref = make(library).parse_reference(url)

    assert ref == SourceRef(
        "file", str(library / "Specs" / "Plan.xlsx"), display_ref="Plan.xlsx", locator=url
    )


def test_the_longest_matching_sync_root_wins(tmp_path: Path) -> None:
    whole, sub = tmp_path / "whole", tmp_path / "sub"
    connector = SharePointConnector(
        config(
            roots_setting(
                (LIBRARY_URL, whole), (LIBRARY_URL + "/Specs", sub), (LIBRARY_URL + "/Spec", tmp_path)
            )
        )
    )

    deep = connector.parse_reference(f"{ENCODED_LIBRARY_URL}/Specs/a.txt")
    shallow = connector.parse_reference(f"{ENCODED_LIBRARY_URL}/Other/a.txt")
    boundary = connector.parse_reference(f"{ENCODED_LIBRARY_URL}/Specsheet/a.txt")

    assert deep is not None and deep.external_id == str(sub / "a.txt")
    assert shallow is not None and shallow.external_id == str(whole / "Other" / "a.txt")
    assert boundary is not None and boundary.external_id == str(whole / "Specsheet" / "a.txt")


def test_root_matching_ignores_case_and_a_trailing_slash(tmp_path: Path) -> None:
    connector = SharePointConnector(
        config(roots_setting(("HTTPS://ACME.SharePoint.com/Sites/hw/shared documents/", tmp_path)))
    )

    ref = connector.parse_reference("https://Acme.sharepoint.com/sites/HW/Shared%20Documents/Specs/A.txt")

    assert ref is not None
    assert ref.external_id == str(tmp_path / "Specs" / "A.txt")


def test_an_unmapped_library_is_rejected_naming_it(library: Path) -> None:
    url = "https://acme.sharepoint.com/sites/Quality/Shared%20Documents/Audits/2026/a.docx?web=1"

    with pytest.raises(ValidationError) as caught:
        make(library).parse_reference(url)

    assert str(caught.value) == (
        "this SharePoint library is not synced on this computer: acme.sharepoint.com/sites/Quality/"
        "Shared Documents — sync it with OneDrive and add it to the connector's sync_roots (CON-51)"
    )


def test_a_url_is_rejected_when_no_sync_root_is_configured() -> None:
    with pytest.raises(ValidationError, match="not synced on this computer"):
        SharePointConnector(config({})).parse_reference(f"{ENCODED_LIBRARY_URL}/a.txt")


def test_a_sharing_prefix_does_not_hide_an_unmapped_library(library: Path) -> None:
    with pytest.raises(ValidationError, match=r"acme\.sharepoint\.com/sites/Quality/Shared Documents"):
        make(library).parse_reference(
            "https://acme.sharepoint.com/:w:/r/sites/Quality/Shared%20Documents/a.docx"
        )


@pytest.mark.parametrize(
    "url",
    [
        "https://acme.sharepoint.com/:x:/s/HW/EabcDEFghi?e=1",
        "https://acme.sharepoint.com/:w:/s/HW",
    ],
)
def test_an_opaque_sharing_link_is_rejected(url: str, library: Path) -> None:
    with pytest.raises(ValidationError, match="opaque SharePoint sharing links cannot be mapped"):
        make(library).parse_reference(url)


@pytest.mark.parametrize(
    "url",
    [
        "https://example.com/sites/HW/Shared%20Documents/a.txt",
        "https://acme.atlassian.net/wiki/spaces/ABC/pages/1",
        "https://sharepoint.com/sites/HW/Shared%20Documents/a.txt",
        "https://acme.sharepoint.com.evil.example/sites/HW/Shared%20Documents/a.txt",
        "http://acme.sharepoint.com/sites/HW/Shared%20Documents/a.txt",
    ],
)
def test_other_urls_are_not_sharepoint_references(url: str, library: Path) -> None:
    assert make(library).parse_reference(url) is None


# --- resolve ---------------------------------------------------------------------------------------------


def test_resolve_an_existing_file(library: Path) -> None:
    path = write(library / "Specs" / "torque.txt")
    url = f"{ENCODED_LIBRARY_URL}/Specs/torque.txt"
    connector = make(library)
    ref = connector.parse_reference(url)
    assert ref is not None

    assert connector.resolve(ref) == ObjectMeta(
        "file", path, title="torque.txt", display_ref="torque.txt", locator=url
    )


def test_resolve_a_missing_file_is_rejected(library: Path) -> None:
    connector = make(library)
    ref = connector.parse_reference(f"{ENCODED_LIBRARY_URL}/Specs/missing.txt")
    assert ref is not None

    with pytest.raises(ValidationError) as caught:
        connector.resolve(ref)

    assert str(caught.value) == (
        f"file not found: {library / 'Specs' / 'missing.txt'} — is the library synced and the file available?"
    )


def test_resolve_rejects_other_object_types(library: Path) -> None:
    with pytest.raises(ValidationError, match="web_link"):
        make(library).resolve(SourceRef("web_link", "https://example.com"))


# --- fetch -----------------------------------------------------------------------------------------------


def test_the_first_fetch_extracts_the_text_and_keeps_no_raw_copy(library: Path, fs: FakeFileSystem) -> None:
    path = str(library / "Specs" / "spec.docx")
    document = docx.Document()
    document.add_paragraph("Torque limit is 12 Nm.")
    document.save(path)
    os.utime(path, (MTIME, MTIME))

    result = fetch(make(library, fs), path)

    version = result.version
    assert result.status == "ok"
    assert version is not None
    assert version.text == "Torque limit is 12 Nm."
    assert version.raw is None
    assert version.media_type is None
    assert version.source_modified_at == "2023-11-14T22:13:20.000Z"
    assert re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d\.\d{3}Z", version.source_modified_at)
    sha256 = version.metadata["sha256"]
    assert isinstance(sha256, str) and len(sha256) == 64
    assert version.source_version_ref == sha256[:12]
    assert version.metadata == {
        "filename": "spec.docx",
        "size": os.path.getsize(path),
        "sha256": sha256,
        "extract_status": "ok",
        "extract_detail": None,
        "truncated": False,
    }
    assert result.title == "spec.docx"
    assert result.warnings == ()
    assert result.entries == () and result.children == ()
    assert result.cursor == (
        f'{{"mtime": {float(MTIME)}, "sha256": "{sha256}", "size": {os.path.getsize(path)}}}'
    )


def test_text_is_capped_at_max_text_chars(library: Path) -> None:
    path = write(library / "long.txt", "x" * 50)

    version = fetch(make(library, max_text_chars=10), path).version

    assert version is not None
    assert (version.text, version.metadata["truncated"]) == ("x" * 10, True)


def test_an_unchanged_mtime_and_size_reads_nothing(library: Path, fs: FakeFileSystem) -> None:
    path = write(library / "a.txt")
    connector = make(library, fs)
    first = fetch(connector, path)
    assert fs.reads == [path]

    again = fetch(connector, path, first.cursor)

    assert fs.reads == [path]
    assert again == FetchResult(cursor=first.cursor, title="a.txt")


def test_a_touched_file_with_the_same_content_gets_a_new_cursor_but_no_version(
    library: Path, fs: FakeFileSystem
) -> None:
    path = write(library / "a.txt")
    connector = make(library, fs)
    first = fetch(connector, path)
    os.utime(path, (MTIME + 60, MTIME + 60))

    again = fetch(connector, path, first.cursor)

    assert fs.reads == [path, path]
    assert again.version is None
    assert again.cursor is not None and again.cursor != first.cursor
    assert f'"mtime": {float(MTIME + 60)}' in again.cursor
    assert again.title == "a.txt"


def test_changed_content_is_a_new_version(library: Path, fs: FakeFileSystem) -> None:
    path = write(library / "a.txt", "Lead time is 26 weeks.")
    connector = make(library, fs)
    first = fetch(connector, path)
    write(library / "a.txt", "Lead time is 30 weeks.", mtime=MTIME + 60)

    again = fetch(connector, path, first.cursor)

    assert again.version is not None
    assert again.version.text == "Lead time is 30 weeks."
    assert again.version.source_modified_at == "2023-11-14T22:14:20.000Z"
    assert first.version is not None
    assert again.version.source_version_ref != first.version.source_version_ref
    assert again.version.raw is None


def test_a_change_with_the_same_mtime_but_another_size_is_noticed(library: Path, fs: FakeFileSystem) -> None:
    path = write(library / "a.txt", "short")
    connector = make(library, fs)
    first = fetch(connector, path)
    write(library / "a.txt", "much longer than before", mtime=MTIME)

    again = fetch(connector, path, first.cursor)

    assert again.version is not None and again.version.text == "much longer than before"


def test_an_offline_placeholder_is_skipped_when_hydrate_is_off(library: Path, fs: FakeFileSystem) -> None:
    path = write(library / "a.txt")
    fs.offline.add(path)

    result = fetch(make(library, fs, hydrate=False), path)

    assert fs.reads == []
    assert result == FetchResult(cursor=None, title="a.txt", warnings=(OFFLINE_SKIPPED,))


def test_a_skipped_placeholder_keeps_the_previous_cursor(library: Path, fs: FakeFileSystem) -> None:
    path = write(library / "a.txt", "old")
    first = fetch(make(library, fs), path)
    write(library / "a.txt", "new content", mtime=MTIME + 60)
    fs.offline.add(path)
    fs.reads.clear()

    result = fetch(make(library, fs, hydrate=False), path, first.cursor)

    assert fs.reads == []
    assert result.version is None
    assert result.cursor == first.cursor
    assert result.warnings == (OFFLINE_SKIPPED,)


def test_an_offline_placeholder_is_hydrated_by_default(library: Path, fs: FakeFileSystem) -> None:
    path = write(library / "a.txt")
    fs.offline.add(path)

    result = fetch(make(library, fs), path)

    assert fs.reads == [path]
    assert result.version is not None and result.version.text == "Lead time is 26 weeks."
    assert result.warnings == (HYDRATED,)


def test_conflict_copies_are_reported_but_not_followed(library: Path, fs: FakeFileSystem) -> None:
    path = write(library / "plan.txt")
    write(library / "plan-LAPTOP01.txt", "other")
    write(library / "plan-PC_2-1.txt", "other")
    write(library / "plan-WORKPC.TXT", "other")
    write(library / "plan-LAPTOP01.docx", "other extension")
    write(library / "planning.txt", "other name")
    write(library / "plan-old copy.txt", "not a computer name")
    write(library / "plan-a-b-c.txt", "too many parts")
    (library / "plan-FOLDER.txt").mkdir()

    result = fetch(make(library, fs), path)

    assert result.warnings == (
        "conflict copy next to this file: plan-LAPTOP01.txt (not followed)",
        "conflict copy next to this file: plan-PC_2-1.txt (not followed)",
        "conflict copy next to this file: plan-WORKPC.TXT (not followed)",
    )
    assert result.version is not None and result.version.text == "Lead time is 26 weeks."
    assert fs.reads == [path]


def test_conflict_copies_are_also_reported_when_the_file_is_unchanged(library: Path) -> None:
    path = write(library / "plan.txt")
    write(library / "plan-LAPTOP01.txt", "other")
    connector = make(library)
    first = fetch(connector, path)

    again = fetch(connector, path, first.cursor)

    assert again.version is None
    assert again.warnings == ("conflict copy next to this file: plan-LAPTOP01.txt (not followed)",)


def test_a_conflict_copy_does_not_flag_the_original_in_turn(library: Path) -> None:
    write(library / "plan.txt")
    copy = write(library / "plan-LAPTOP01.txt", "other")

    assert fetch(make(library), copy).warnings == ()


def test_a_deleted_or_moved_file_is_gone(library: Path) -> None:
    path = write(library / "a.txt")
    os.rename(path, library / "b.txt")

    assert fetch(make(library), path) == FetchResult(status="gone")


def test_an_unsupported_extension_is_reported_as_a_warning(library: Path) -> None:
    path = write(library / "drawing.dwg", "binary")

    result = fetch(make(library), path)

    assert result.version is not None
    assert result.version.text == ""
    assert result.version.metadata["extract_status"] == "unsupported"
    assert result.warnings == ("text extraction: unsupported — no extractor for .dwg",)


def test_a_broken_document_is_reported_as_a_warning(library: Path) -> None:
    path = write(library / "broken.docx", "this is not a zip file")

    result = fetch(make(library), path)

    assert result.version is not None
    assert result.version.metadata["extract_status"] == "error"
    assert len(result.warnings) == 1 and result.warnings[0].startswith("text extraction: error — ")


# --- check -----------------------------------------------------------------------------------------------


def test_check_reports_every_sync_root(library: Path, tmp_path: Path) -> None:
    missing = tmp_path / "not synced"
    connector = SharePointConnector(
        config(roots_setting((LIBRARY_URL, library), (LIBRARY_URL + "-2", missing)))
    )

    report = connector.check()

    assert [(c.name, c.ok, c.detail) for c in report.checks] == [
        (f"sync root {LIBRARY_URL}", True, str(library)),
        (f"sync root {LIBRARY_URL}-2", False, f"folder not found: {missing}"),
    ]
    assert not report.ok


def test_check_with_existing_roots_is_ok(library: Path) -> None:
    assert make(library).check().ok


def test_check_without_sync_roots_fails(library: Path) -> None:
    report = SharePointConnector(config({})).check()

    assert [(c.name, c.ok, c.detail) for c in report.checks] == [
        ("sync roots", False, "no sync_roots configured")
    ]


# --- search_candidates -----------------------------------------------------------------------------------


@pytest.fixture
def tree(library: Path, tmp_path: Path) -> tuple[SharePointConnector, dict[str, str]]:
    other = tmp_path / "OneDrive" / "Personal"
    other.mkdir()
    paths = {
        "old": write(library / "Torque report v1.txt", mtime=MTIME),
        "mid": write(library / "Specs" / "TORQUE Report v2.txt", mtime=MTIME + 100),
        "new": write(other / "my torque report.txt", mtime=MTIME + 200),
        "unrelated": write(library / "Specs" / "bom.txt", mtime=MTIME + 300),
    }
    connector = SharePointConnector(
        config(roots_setting((LIBRARY_URL, library), ("https://acme-my.sharepoint.com/personal/alex", other)))
    )
    return connector, paths


def meta(path: str) -> ObjectMeta:
    name = os.path.basename(path)
    return ObjectMeta("file", path, title=name, display_ref=name, locator=path)


def test_search_matches_names_newest_first(tree) -> None:
    connector, paths = tree

    found = connector.search_candidates(CandidateQuery("TORQUE Report"))

    assert found == [meta(paths["new"]), meta(paths["mid"]), meta(paths["old"])]


def test_search_without_a_match(tree) -> None:
    connector, _ = tree

    assert connector.search_candidates(CandidateQuery("nothing like this")) == []


def test_search_filters_by_modification_time(tree) -> None:
    connector, paths = tree

    found = connector.search_candidates(CandidateQuery("report", since="2023-11-14T22:14:00.000Z"))

    assert found == [meta(paths["new"]), meta(paths["mid"])]


def test_search_since_is_inclusive(tree) -> None:
    connector, paths = tree

    found = connector.search_candidates(CandidateQuery("report", since="2023-11-14T22:14:20.000Z"))

    assert [item.external_id for item in found] == [paths["new"], paths["mid"]]


def test_search_is_limited_to_the_newest(tree) -> None:
    connector, paths = tree

    assert connector.search_candidates(CandidateQuery("report", limit=2)) == [
        meta(paths["new"]),
        meta(paths["mid"]),
    ]


def test_search_lists_a_file_of_overlapping_roots_once(library: Path) -> None:
    path = write(library / "Specs" / "torque.txt")
    connector = SharePointConnector(
        config(roots_setting((LIBRARY_URL, library), (LIBRARY_URL + "/Specs", library / "Specs")))
    )

    assert connector.search_candidates(CandidateQuery("torque")) == [meta(path)]


# --- end to end through the store and the sync engine ----------------------------------------------------


@pytest.fixture
def linked(engine: sqlalchemy.Engine, tmp_path: Path, library: Path):
    home = tmp_path / "sync_home"
    write_default_config(home)
    path = write(library / "Specs" / "torque.txt", "Torque limit is 12 Nm.")
    settings = roots_setting((LIBRARY_URL, library))
    with engine.begin() as conn:
        me = people.create_person(conn, "Alex", is_me=True)
        projects.create_project(conn, "OIL", "Open issues")
        issue = issues.create_issue(conn, "OIL", "Inductor delivery", "Chase the supplier", actor_id=me.id)
        row = registry.create_connector(conn, "sharepoint", "SharePoint", settings=settings)
        connector = SharePointConnector(config(settings))
        ref = connector.parse_reference(f"{ENCODED_LIBRARY_URL}/Specs/torque.txt")
        assert ref is not None
        object_id = store.upsert_object(conn, row.id, connector.resolve(ref))
        linking.link_source(conn, issue.key, object_id, actor_id=me.id)

    def sync():
        return run_sync(
            engine,
            SyncScope("source", object_id),
            config=load_config(home),
            connector_factory=lambda connector_row: SharePointConnector(config(settings)),
        )

    return engine, object_id, path, sync


def test_the_first_sync_stores_one_version_without_a_blob(linked) -> None:
    engine, object_id, path, sync = linked

    report = sync()

    assert report.status == "ok"
    assert (report.stats["new_versions"], report.stats["new_entries"]) == (1, 0)
    with engine.connect() as conn:
        obj = store.get_object(conn, object_id)
        [version] = store.list_versions(conn, object_id)
    assert (obj.object_type, obj.external_id, obj.display_ref, obj.title, obj.fetch_status) == (
        "file",
        path,
        "torque.txt",
        "torque.txt",
        "ok",
    )
    assert obj.locator == f"{ENCODED_LIBRARY_URL}/Specs/torque.txt"
    assert obj.cursor is not None and '"sha256"' in obj.cursor
    assert version.text == "Torque limit is 12 Nm."
    assert version.source_modified_at == "2023-11-14T22:13:20.000Z"


def test_versions_are_stored_only_when_the_content_changes(linked, library: Path) -> None:
    engine, object_id, path, sync = linked
    sync()
    unchanged = sync()
    os.utime(path, (MTIME + 60, MTIME + 60))
    touched = sync()
    write(library / "Specs" / "torque.txt", "Torque limit is 15 Nm.", mtime=MTIME + 120)
    edited = sync()

    assert [r.stats["new_versions"] for r in (unchanged, touched, edited)] == [0, 0, 1]
    assert [r.results[0].outcome for r in (unchanged, touched, edited)] == [
        "unchanged",
        "unchanged",
        "changed",
    ]
    with engine.connect() as conn:
        assert [v.text for v in store.list_versions(conn, object_id)] == [
            "Torque limit is 12 Nm.",
            "Torque limit is 15 Nm.",
        ]
        cursor = store.get_object(conn, object_id).cursor
    assert cursor is not None and f'"mtime": {float(MTIME + 120)}' in cursor


def test_a_conflict_copy_becomes_a_warning_in_the_sync_report(linked, library: Path) -> None:
    _, _, _, sync = linked
    write(library / "Specs" / "torque-LAPTOP01.txt", "conflict")

    report = sync()

    assert report.results[0].warnings == (
        "conflict copy next to this file: torque-LAPTOP01.txt (not followed)",
    )


def test_a_moved_file_becomes_gone_and_keeps_its_versions(linked) -> None:
    engine, object_id, path, sync = linked
    sync()
    os.rename(path, path + ".moved")

    report = sync()

    assert report.results[0].outcome == "gone"
    with engine.connect() as conn:
        assert store.get_object(conn, object_id).fetch_status == "gone"
        assert len(store.list_versions(conn, object_id)) == 1


# --- registration ----------------------------------------------------------------------------------------


def test_the_connector_is_registered_through_its_entry_point() -> None:
    assert registry.available_connector_types()["sharepoint"] is SharePointConnector
