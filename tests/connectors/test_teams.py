import json
import os
from datetime import UTC, datetime
from pathlib import Path

import pytest
import sqlalchemy
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from ai_oil import timeutil
from ai_oil.config import load_config, write_default_config
from ai_oil.connectors.sharepoint import FileInfo, LocalFileSystem
from ai_oil.connectors.teams import ConversationCandidate, TeamsConnector
from ai_oil.domain import issues, people, projects
from ai_oil.errors import ValidationError
from ai_oil.search import fts
from ai_oil.sources import linking, registry, store
from ai_oil.sources.model import (
    CandidateQuery,
    ConnectorConfig,
    FetchResult,
    IdentityRef,
    ObjectMeta,
    ObjectState,
    SourceRef,
)
from ai_oil.sources.sync import SyncScope, run_sync

MTIME = 1_700_000_000
COLUMNS = (
    "MessageId",
    "ConversationId",
    "ConversationType",
    "ConversationName",
    "Timestamp",
    "SenderName",
    "SenderEmail",
    "Permalink",
    "Text",
    "Trigger",
)
CHAT = "19:chat-one@thread.v2"
CHAT_ENCODED = "19%3Achat-one%40thread.v2"
OTHER_CHAT = "19:chat-two@thread.v2"
CHANNEL = "19:channel@thread.tacv2"
LINK = "https://teams.microsoft.com/l/message"
OFFLINE_SKIPPED = "not available offline — skipped (hydrate is off)"


def at(day: int, hour: int = 10, minute: int = 0) -> datetime:
    return datetime(2026, 3, day, hour, minute, tzinfo=UTC)


def msg(message_id: str, **fields: object) -> dict[str, object]:
    values: dict[str, object] = {
        "MessageId": message_id,
        "ConversationId": CHAT,
        "ConversationType": "chat",
        "ConversationName": "",
        "Timestamp": "2026-03-01T10:00:00+00:00",
        "SenderName": "Jane Doe",
        "SenderEmail": "Jane.Doe@acme.com",
        "Permalink": "",
        "Text": "<p>Hello</p>",
        "Trigger": "Chat",
    }
    values.update(fields)
    return values


class FakeFileSystem:
    """The local file system with a marker for Files On-Demand placeholders and a log of the reads."""

    def __init__(self) -> None:
        self._local = LocalFileSystem()
        self.offline: set[str] = set()
        self.reads: list[str] = []

    def is_file(self, path: str) -> bool:
        return self._local.is_file(path)

    def is_dir(self, path: str) -> bool:
        return self._local.is_dir(path)

    def info(self, path: str) -> FileInfo:
        real = self._local.info(path)
        return FileInfo(real.mtime, real.size, path in self.offline)

    def read_bytes(self, path: str) -> bytes:
        self.reads.append(os.path.basename(path))
        return self._local.read_bytes(path)

    def list_dir(self, path: str) -> list[str]:
        return self._local.list_dir(path)

    def walk_files(self, root: str):
        return self._local.walk_files(root)


def config(settings: dict) -> ConnectorConfig:
    return ConnectorConfig(1, "teams", "Teams", True, settings, None)


@pytest.fixture
def folder(tmp_path: Path) -> Path:
    path = tmp_path / "AI_OIL-Teams"
    path.mkdir()
    return path


@pytest.fixture
def fs() -> FakeFileSystem:
    return FakeFileSystem()


def make(folder: Path, fs: FakeFileSystem | None = None, **settings: object) -> TeamsConnector:
    return TeamsConnector(config({"folder": str(folder), **settings}), fs=fs)


def bridge(
    folder: Path,
    rows: list[dict[str, object]],
    name: str = "Teams_Bridge_2026.xlsx",
    *,
    columns: tuple[str, ...] = COLUMNS,
    table: str = "TeamsMessages",
    sheet: str = "Messages",
    mtime: int = MTIME,
) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(list(columns))
    for row in rows:
        worksheet.append([row.get(column) for column in columns])
    last = f"{get_column_letter(len(columns))}{len(rows) + 1}"
    worksheet.add_table(Table(displayName=table, ref=f"A1:{last}"))
    path = folder / name
    workbook.save(path)
    os.utime(path, (mtime, mtime))
    return path


def state(conversation_id: str = CHAT, cursor: str | None = None) -> ObjectState:
    return ObjectState(
        id=1,
        connector_id=1,
        object_type="teams_thread",
        external_id=conversation_id,
        title=None,
        display_ref=None,
        locator=None,
        cursor=cursor,
        parent_id=None,
        latest_content_hash=None,
    )


def fetch(connector: TeamsConnector, conversation_id: str = CHAT) -> FetchResult:
    return connector.fetch(state(conversation_id))


# --- settings and contract -------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("settings", "key"),
    [
        ({}, "folder"),
        ({"folder": " "}, "folder"),
        ({"folder": 5}, "folder"),
        ({"folder": "/x", "file_pattern": ""}, "file_pattern"),
        ({"folder": "/x", "file_pattern": ["a"]}, "file_pattern"),
        ({"folder": "/x", "table": " "}, "table"),
        ({"folder": "/x", "table": 5}, "table"),
        ({"folder": "/x", "hydrate": "yes"}, "hydrate"),
    ],
)
def test_invalid_settings_are_rejected_naming_the_key(settings: dict, key: str) -> None:
    with pytest.raises(ValidationError, match=key):
        TeamsConnector(config(settings))


def test_the_defaults_need_only_the_folder() -> None:
    connector = TeamsConnector(config({"folder": " C:/Users/alex/Documents/AI_OIL-Teams "}))

    assert isinstance(connector._fs, LocalFileSystem)
    assert connector._folder == "C:/Users/alex/Documents/AI_OIL-Teams"
    assert connector._pattern == "Teams_Bridge*.xlsx"
    assert connector._table == "TeamsMessages"
    assert connector._hydrate is True


def test_the_connector_declares_its_contract() -> None:
    assert TeamsConnector.type == "teams"
    assert TeamsConnector.object_types == frozenset({"teams_thread"})
    assert TeamsConnector.action_class == "A"
    assert TeamsConnector.max_parallel == 1


def test_the_plugin_is_registered_under_its_type_name() -> None:
    assert registry.get_connector_class("teams") is TeamsConnector


# --- reading the bridge files ----------------------------------------------------------------------------


def test_columns_are_read_by_name_in_any_order_and_unknown_columns_are_ignored(folder: Path) -> None:
    shuffled = ("Text", "Extra", "Timestamp", "SenderName", "ConversationType", "ConversationId", "MessageId")
    bridge(
        folder,
        [msg("m1", Text="Torque is 12 Nm", Extra="ignored")],
        columns=shuffled,
    )

    result = fetch(make(folder))

    assert result.status == "ok" and result.warnings == ()
    [entry] = result.entries
    assert (entry.external_id, entry.text, entry.author) == ("m1", "Torque is 12 Nm", None)


def test_header_names_ignore_case_and_the_table_name_ignores_case(folder: Path) -> None:
    bridge(
        folder,
        [{key.lower(): value for key, value in msg("m1").items()}],
        columns=tuple(column.lower() for column in COLUMNS),
        table="teamsmessages",
    )

    assert fetch(make(folder)).warnings == ()


def test_the_table_is_found_on_any_sheet(folder: Path) -> None:
    bridge(folder, [msg("m1")], sheet="Sheet7")

    assert [e.external_id for e in fetch(make(folder)).entries] == ["m1"]


def test_a_custom_table_and_file_pattern(folder: Path) -> None:
    bridge(folder, [msg("m1")], name="Chats_2026.xlsx", table="Chats")
    bridge(folder, [msg("m2")], name="Teams_Bridge_2026.xlsx")

    result = fetch(make(folder, file_pattern="chats_*.xlsx", table="CHATS"))

    assert [e.external_id for e in result.entries] == ["m1"]


def test_a_file_without_the_table_is_a_warning_and_other_files_are_still_read(folder: Path) -> None:
    bridge(folder, [msg("m1")], name="Teams_Bridge_2025.xlsx", table="Other")
    bridge(folder, [msg("m2", Timestamp="2026-03-02T10:00:00+00:00")], name="Teams_Bridge_2026.xlsx")

    result = fetch(make(folder))

    assert [e.external_id for e in result.entries] == ["m2"]
    assert result.warnings == ("Teams_Bridge_2025.xlsx: table TeamsMessages not found",)


def test_a_missing_required_column_is_a_warning_naming_file_and_column(folder: Path) -> None:
    trimmed = tuple(column for column in COLUMNS if column not in ("Timestamp", "SenderName"))
    bridge(folder, [msg("m1")], name="Teams_Bridge_2025.xlsx", columns=trimmed)
    bridge(folder, [msg("m2")], name="Teams_Bridge_2026.xlsx")

    result = fetch(make(folder))

    assert [e.external_id for e in result.entries] == ["m2"]
    assert result.warnings == (
        "Teams_Bridge_2025.xlsx: table TeamsMessages lacks the required column(s): Timestamp, SenderName",
    )


def test_optional_columns_may_be_absent(folder: Path) -> None:
    required = ("MessageId", "ConversationId", "ConversationType", "Timestamp", "SenderName", "Text")
    bridge(folder, [msg("m1")], columns=required)

    result = fetch(make(folder))

    assert result.warnings == ()
    assert result.title == "Chat with Jane Doe"
    [entry] = result.entries
    assert entry.author is None and entry.metadata == {"sender_name": "Jane Doe"}


def test_a_file_that_is_not_a_workbook_is_a_warning(folder: Path) -> None:
    (folder / "Teams_Bridge_2025.xlsx").write_text("not a workbook", encoding="utf-8")
    bridge(folder, [msg("m1")])

    result = fetch(make(folder))

    assert [e.external_id for e in result.entries] == ["m1"]
    [warning] = result.warnings
    assert warning.startswith("Teams_Bridge_2025.xlsx: cannot read the workbook: ")


def test_files_are_read_in_name_order_and_other_files_are_not_touched(folder: Path, fs) -> None:
    bridge(folder, [msg("m2")], name="Teams_Bridge_2027.xlsx")
    bridge(folder, [msg("m1")], name="Teams_Bridge_2026.xlsx")
    (folder / "notes.txt").write_text("x", encoding="utf-8")
    (folder / "~$Teams_Bridge_2026.xlsx").write_bytes(b"lock")
    (folder / "Teams_Bridge_old").mkdir()

    fetch(make(folder, fs))

    assert fs.reads == ["Teams_Bridge_2026.xlsx", "Teams_Bridge_2027.xlsx"]


def test_a_message_written_by_both_flows_is_stored_once(folder: Path) -> None:
    bridge(
        folder,
        [msg("m1", Trigger="Mention", ConversationName="", SenderEmail="", Permalink="")],
        name="Teams_Bridge_2026.xlsx",
    )
    bridge(
        folder,
        [
            msg("m1", Trigger="Chat", ConversationName="Torque chat", Permalink="https://t/m1"),
            msg("m2", Timestamp="2026-03-01T11:00:00+00:00"),
        ],
        name="Teams_Bridge_2027.xlsx",
    )

    result = fetch(make(folder))

    assert [e.external_id for e in result.entries] == ["m1", "m2"]
    assert result.title == "Torque chat"
    first = result.entries[0]
    assert first.author == IdentityRef("email", "jane.doe@acme.com", "Jane Doe")
    assert first.metadata["permalink"] == "https://t/m1"


def test_a_message_repeated_within_one_file_is_stored_once(folder: Path) -> None:
    bridge(folder, [msg("m1"), msg("m1", Trigger="Mention")])

    assert [e.external_id for e in fetch(make(folder)).entries] == ["m1"]


def test_iso_timestamps_with_any_offset_become_utc(folder: Path) -> None:
    bridge(
        folder,
        [
            msg("m1", Timestamp="2026-03-01T11:00:00+01:00"),
            msg("m2", Timestamp="2026-03-01T10:30:00.1234567Z"),
            msg("m3", Timestamp="2026-03-01T12:00:00Z"),
        ],
    )

    result = fetch(make(folder))

    assert [(e.external_id, e.occurred_at) for e in result.entries] == [
        ("m1", "2026-03-01T10:00:00.000Z"),
        ("m2", "2026-03-01T10:30:00.123Z"),
        ("m3", "2026-03-01T12:00:00.000Z"),
    ]
    assert result.warnings == ()


def test_naive_timestamps_are_taken_as_utc_with_one_warning_per_file(folder: Path) -> None:
    bridge(
        folder,
        [
            msg("m1", Timestamp=datetime(2026, 3, 1, 10, 0)),
            msg("m2", Timestamp="2026-03-01T11:00:00"),
            msg("m3", Timestamp="2026-03-01T12:00:00+00:00"),
        ],
    )

    result = fetch(make(folder))

    assert [e.occurred_at for e in result.entries] == [
        "2026-03-01T10:00:00.000Z",
        "2026-03-01T11:00:00.000Z",
        "2026-03-01T12:00:00.000Z",
    ]
    assert result.warnings == ("Teams_Bridge_2026.xlsx: timestamps without a time zone were taken as UTC",)


def test_rows_without_the_essentials_are_skipped_and_counted(folder: Path) -> None:
    bridge(
        folder,
        [
            msg("m1"),
            msg(""),
            msg("m3", ConversationId=""),
            msg("m4", Timestamp="yesterday"),
            msg("m5", Timestamp=None),
            msg("m6", Timestamp=45000),
        ],
    )

    result = fetch(make(folder))

    assert [e.external_id for e in result.entries] == ["m1"]
    assert result.warnings == (
        "Teams_Bridge_2026.xlsx: 5 row(s) skipped: no MessageId, no ConversationId or no valid Timestamp",
    )


def test_numeric_ids_are_read_as_text(folder: Path) -> None:
    bridge(folder, [msg(1700000000000, ConversationId=19.0)])

    result = fetch(make(folder), "19")

    assert [e.external_id for e in result.entries] == ["1700000000000"]


# --- fetch -----------------------------------------------------------------------------------------------


def test_fetch_returns_time_ordered_entries_and_the_newest_time_as_cursor_but_no_version(
    folder: Path,
) -> None:
    bridge(
        folder,
        [
            msg(
                "m2",
                Timestamp="2026-03-01T11:30:00+00:00",
                SenderName="Alex",
                SenderEmail="alex@example.com",
                Text="<p>Will do.</p><p>Friday.</p>",
            ),
            msg("m1", Timestamp="2026-03-01T10:00:00+00:00", Text="<p>Lead time &amp; price?</p>"),
            msg("o1", ConversationId=OTHER_CHAT, Text="elsewhere"),
        ],
    )

    result = fetch(make(folder))

    assert result.status == "ok"
    assert [(e.entry_type, e.external_id, e.occurred_at) for e in result.entries] == [
        ("teams_message", "m1", "2026-03-01T10:00:00.000Z"),
        ("teams_message", "m2", "2026-03-01T11:30:00.000Z"),
    ]
    assert [e.text for e in result.entries] == ["Lead time & price?", "Will do.\nFriday."]
    assert result.entries[0].new_text is None and result.entries[0].direction is None
    assert result.version is None
    assert json.loads(result.cursor or "") == {"latest": "2026-03-01T11:30:00.000Z"}
    assert result.children == ()


def test_the_author_is_the_lower_cased_sender_email_with_the_display_name(folder: Path) -> None:
    bridge(
        folder,
        [
            msg("m1", SenderEmail="Jane.Doe@ACME.com"),
            msg("m2", SenderEmail="", SenderName="Guest"),
            msg("m3", SenderEmail="jane-upn", SenderName="Odd"),
        ],
    )

    first, second, third = fetch(make(folder)).entries

    assert first.author == IdentityRef("email", "jane.doe@acme.com", "Jane Doe")
    assert second.author is None and second.metadata == {"sender_name": "Guest"}
    assert third.author is None and third.metadata == {"sender_name": "Odd", "sender_raw": "jane-upn"}


def test_the_permalink_goes_into_the_entry_metadata(folder: Path) -> None:
    bridge(folder, [msg("m1", Permalink="https://teams.microsoft.com/l/message/x/1")])

    [entry] = fetch(make(folder)).entries

    assert entry.metadata == {
        "sender_name": "Jane Doe",
        "permalink": "https://teams.microsoft.com/l/message/x/1",
    }


def test_the_title_is_the_newest_conversation_name(folder: Path) -> None:
    bridge(
        folder,
        [
            msg("m1", ConversationName="Old topic"),
            msg("m2", ConversationName="", Timestamp="2026-03-01T11:00:00+00:00"),
            msg("m3", ConversationName="Torque topic", Timestamp="2026-03-01T12:00:00+00:00"),
            msg("m4", ConversationName="", Timestamp="2026-03-01T13:00:00+00:00"),
        ],
    )

    assert fetch(make(folder)).title == "Torque topic"


@pytest.mark.parametrize(
    ("senders", "title"),
    [
        (["Jane"], "Chat with Jane"),
        (["Jane", "Bob", "Jane"], "Chat with Jane, Bob"),
        (["A", "B", "C", "D"], "Chat with A, B, C, D"),
        (["A", "B", "C", "D", "E", "F", "A"], "Chat with A, B, C, D and 2 more"),
        ([""], "Chat"),
    ],
)
def test_an_unnamed_chat_is_titled_after_its_senders(folder: Path, senders: list[str], title: str) -> None:
    rows = [
        msg(f"m{index}", SenderName=name, Timestamp=f"2026-03-01T10:{index:02d}:00+00:00")
        for index, name in enumerate(senders)
    ]
    bridge(folder, rows)

    assert fetch(make(folder)).title == title


def test_an_unnamed_channel_conversation_has_a_fixed_title(folder: Path) -> None:
    bridge(folder, [msg("m1", ConversationId=CHANNEL, ConversationType="Channel")])

    assert fetch(make(folder), CHANNEL).title == "Channel conversation"


def test_a_named_channel_keeps_its_name(folder: Path) -> None:
    bridge(
        folder, [msg("m1", ConversationId=CHANNEL, ConversationType="channel", ConversationName="MB32829")]
    )

    assert fetch(make(folder), CHANNEL).title == "MB32829"


def test_an_unchanged_refetch_gives_the_same_entries_and_a_grown_conversation_only_adds_one(
    folder: Path,
) -> None:
    rows = [msg("m1"), msg("m2", Timestamp="2026-03-01T11:00:00+00:00", Trigger="Mention")]
    bridge(folder, rows)
    first = fetch(make(folder))
    bridge(folder, list(reversed(rows)), name="Teams_Bridge_2026.xlsx", mtime=MTIME + 60)
    again = fetch(make(folder))
    bridge(folder, [msg("m3", Timestamp="2026-03-01T12:00:00+00:00")], name="Teams_Bridge_2027.xlsx")
    grown = fetch(make(folder))

    hashes = [store.entry_hash(e) for e in first.entries]
    assert hashes == [store.entry_hash(e) for e in again.entries]
    assert [store.entry_hash(e) for e in grown.entries][:2] == hashes
    assert [e.external_id for e in grown.entries] == ["m1", "m2", "m3"]


def test_the_parsed_files_are_cached_per_instance_by_path_mtime_and_size(folder: Path, fs) -> None:
    bridge(folder, [msg("m1"), msg("o1", ConversationId=OTHER_CHAT)])
    connector = make(folder, fs)

    fetch(connector)
    fetch(connector, OTHER_CHAT)
    connector.find_conversations("hello")
    assert fs.reads == ["Teams_Bridge_2026.xlsx"]

    bridge(folder, [msg("m1"), msg("m2")], mtime=MTIME + 60)
    assert [e.external_id for e in fetch(connector).entries] == ["m1", "m2"]
    assert fs.reads == ["Teams_Bridge_2026.xlsx", "Teams_Bridge_2026.xlsx"]
    fetch(make(folder, fs))
    assert len(fs.reads) == 3


def test_an_empty_text_and_a_plain_text_message_are_kept(folder: Path) -> None:
    bridge(folder, [msg("m1", Text=""), msg("m2", Text="a < b", Timestamp="2026-03-01T11:00:00+00:00")])

    first, second = fetch(make(folder)).entries

    assert (first.text, second.text) == ("", "a < b")


# --- bridge-level warnings are reported once per instance -------------------------------------------------


def test_conflict_copies_are_reported_by_the_first_fetch_only(folder: Path) -> None:
    bridge(folder, [msg("m1"), msg("o1", ConversationId=OTHER_CHAT)])
    bridge(folder, [msg("m9")], name="Teams_Bridge_2026-LAPTOP7.xlsx")
    connector = make(folder)

    first = fetch(connector)
    second = fetch(connector, OTHER_CHAT)
    third = fetch(connector)
    fresh = fetch(make(folder))

    assert first.warnings == ("conflict copy of a bridge file: Teams_Bridge_2026-LAPTOP7.xlsx (not read)",)
    assert second.warnings == third.warnings == ()
    assert fresh.warnings == first.warnings


def test_offline_notes_skipped_rows_and_naive_timestamps_are_reported_once(folder: Path, fs) -> None:
    hydrated = bridge(folder, [msg("m1"), msg("m2", Timestamp="nope")], name="Teams_Bridge_2025.xlsx")
    skipped = bridge(folder, [msg("m3")], name="Teams_Bridge_2026.xlsx")
    bridge(folder, [msg("m4", Timestamp=datetime(2026, 3, 1))], name="Teams_Bridge_2027.xlsx")
    fs.offline.update({str(hydrated), str(skipped)})
    connector = make(folder, fs)

    first = fetch(connector)
    second = fetch(connector)

    assert first.warnings == (
        "Teams_Bridge_2025.xlsx: hydrated: downloaded on read (Files On-Demand)",
        "Teams_Bridge_2025.xlsx: 1 row(s) skipped: no MessageId, no ConversationId or no valid Timestamp",
        "Teams_Bridge_2026.xlsx: hydrated: downloaded on read (Files On-Demand)",
        "Teams_Bridge_2027.xlsx: timestamps without a time zone were taken as UTC",
    )
    assert second.warnings == ()


def test_a_skipped_offline_file_is_reported_once(folder: Path, fs) -> None:
    path = bridge(folder, [msg("m1")])
    fs.offline.add(str(path))
    connector = make(folder, fs, hydrate=False)

    first = fetch(connector)
    second = fetch(connector)

    assert first.warnings[0] == f"Teams_Bridge_2026.xlsx: {OFFLINE_SKIPPED}"
    assert second.warnings == (
        "conversation not found, but not every bridge file could be read — not marked as gone",
    )


def test_a_file_that_cannot_be_used_is_reported_on_every_fetch(folder: Path) -> None:
    bridge(folder, [msg("m1")], name="Teams_Bridge_2025.xlsx", table="Other")
    bridge(folder, [msg("m2")], name="Teams_Bridge_2026.xlsx")
    connector = make(folder)

    first = fetch(connector)
    second = fetch(connector)

    assert first.warnings == second.warnings == ("Teams_Bridge_2025.xlsx: table TeamsMessages not found",)


def test_a_check_reports_the_notices_and_the_following_fetch_does_not_repeat_them(folder: Path) -> None:
    bridge(folder, [msg("m1", Timestamp=datetime(2026, 3, 1))])
    bridge(folder, [msg("m9")], name="Teams_Bridge_2026-LAPTOP7.xlsx")
    connector = make(folder)

    checks = {check.name: check for check in connector.check().checks}
    again = {check.name: check for check in connector.check().checks}

    assert "timestamps without a time zone" in checks["file Teams_Bridge_2026.xlsx"].detail
    assert "conflict copy Teams_Bridge_2026-LAPTOP7.xlsx" in checks
    assert checks.keys() == again.keys()
    assert fetch(connector).warnings == ()


def test_a_fetch_before_a_check_uses_up_the_notices_for_the_fetch(folder: Path) -> None:
    bridge(folder, [msg("m1")])
    bridge(folder, [msg("m9")], name="Teams_Bridge_2026-LAPTOP7.xlsx")
    connector = make(folder)

    assert len(fetch(connector).warnings) == 1
    assert "conflict copy Teams_Bridge_2026-LAPTOP7.xlsx" in [c.name for c in connector.check().checks]


# --- gone, incomplete and unreachable ---------------------------------------------------------------------


def test_a_conversation_that_is_in_no_file_is_gone(folder: Path) -> None:
    bridge(folder, [msg("m1")])

    result = fetch(make(folder), OTHER_CHAT)

    assert result.status == "gone"
    assert result.warnings == ("conversation not found in the bridge files (not captured yet, or removed)",)


def test_a_conversation_is_not_declared_gone_while_a_file_could_not_be_read(folder: Path, fs) -> None:
    path = bridge(folder, [msg("m1")], name="Teams_Bridge_2025.xlsx")
    bridge(folder, [msg("m2", ConversationId=OTHER_CHAT)], name="Teams_Bridge_2026.xlsx")
    fs.offline.add(str(path))

    result = fetch(make(folder, fs, hydrate=False))

    assert result.status == "ok" and result.entries == () and result.version is None
    assert result.warnings == (
        f"Teams_Bridge_2025.xlsx: {OFFLINE_SKIPPED}",
        "conversation not found, but not every bridge file could be read — not marked as gone",
    )


def test_an_unreadable_file_also_blocks_gone(folder: Path) -> None:
    (folder / "Teams_Bridge_2025.xlsx").write_text("broken", encoding="utf-8")

    assert fetch(make(folder)).status == "ok"


def test_a_missing_folder_is_an_error_not_gone(tmp_path: Path) -> None:
    with pytest.raises(ValidationError, match="Teams bridge folder not found"):
        fetch(make(tmp_path / "missing"))


# --- offline placeholders and conflict copies ------------------------------------------------------------


def test_an_offline_file_is_hydrated_by_default_and_reported(folder: Path, fs) -> None:
    path = bridge(folder, [msg("m1")])
    fs.offline.add(str(path))

    result = fetch(make(folder, fs))

    assert [e.external_id for e in result.entries] == ["m1"]
    assert result.warnings == ("Teams_Bridge_2026.xlsx: hydrated: downloaded on read (Files On-Demand)",)
    assert fs.reads == ["Teams_Bridge_2026.xlsx"]


def test_an_offline_file_is_skipped_when_hydrate_is_off(folder: Path, fs) -> None:
    path = bridge(folder, [msg("m1")], name="Teams_Bridge_2025.xlsx")
    bridge(folder, [msg("m2", Timestamp="2026-03-01T11:00:00+00:00")])
    fs.offline.add(str(path))

    result = fetch(make(folder, fs, hydrate=False))

    assert [e.external_id for e in result.entries] == ["m2"]
    assert result.warnings == (f"Teams_Bridge_2025.xlsx: {OFFLINE_SKIPPED}",)
    assert fs.reads == ["Teams_Bridge_2026.xlsx"]


def test_a_conflict_copy_is_reported_and_not_read(folder: Path, fs) -> None:
    bridge(folder, [msg("m1")])
    bridge(folder, [msg("m9", Text="from the copy")], name="Teams_Bridge_2026-LAPTOP7.xlsx")
    bridge(folder, [msg("m8", Text="second copy")], name="Teams_Bridge_2026-LAPTOP7-2.xlsx")
    bridge(folder, [msg("m2", Timestamp="2026-03-01T11:00:00+00:00")], name="Teams_Bridge_2027.xlsx")

    result = fetch(make(folder, fs))

    assert [e.external_id for e in result.entries] == ["m1", "m2"]
    assert result.warnings == (
        "conflict copy of a bridge file: Teams_Bridge_2026-LAPTOP7-2.xlsx (not read)",
        "conflict copy of a bridge file: Teams_Bridge_2026-LAPTOP7.xlsx (not read)",
    )
    assert fs.reads == ["Teams_Bridge_2026.xlsx", "Teams_Bridge_2027.xlsx"]


def test_a_file_that_merely_looks_like_a_copy_is_read_when_its_base_is_absent(folder: Path) -> None:
    bridge(folder, [msg("m1")], name="Teams_Bridge_2026-archive.xlsx")

    result = fetch(make(folder))

    assert [e.external_id for e in result.entries] == ["m1"] and result.warnings == ()


# --- parse_reference and resolve -------------------------------------------------------------------------


def test_a_chat_link_is_resolved_through_the_message_id(folder: Path) -> None:
    bridge(
        folder,
        [
            msg("1700000000001", ConversationId=CHAT, ConversationName="Torque chat"),
            msg("1700000000002", ConversationId=OTHER_CHAT, Timestamp="2026-03-01T11:00:00+00:00"),
        ],
    )
    url = f"{LINK}/{CHAT_ENCODED}/1700000000002?tenantId=t&groupId=g&parentMessageId=1&createdTime=1"

    ref = make(folder).parse_reference(f"  {url}\n")

    assert ref == SourceRef("teams_thread", OTHER_CHAT, display_ref="Chat with Jane Doe", locator=url)


def test_a_chat_link_falls_back_to_the_decoded_conversation_id(folder: Path) -> None:
    bridge(folder, [msg("1700000000001", ConversationName="Torque chat")])
    url = f"{LINK}/{CHAT_ENCODED}/1700000000099?tenantId=t"

    ref = make(folder).parse_reference(url)

    assert ref == SourceRef("teams_thread", CHAT, display_ref="Torque chat", locator=url)


def test_a_link_to_a_conversation_that_is_not_captured_has_no_title_yet(folder: Path) -> None:
    bridge(folder, [msg("m1")])
    url = f"{LINK}/19%3Aunseen%40thread.v2/1700000000099"

    ref = make(folder).parse_reference(url)

    assert ref == SourceRef("teams_thread", "19:unseen@thread.v2", display_ref=None, locator=url)


def test_a_channel_link_with_a_parent_message(folder: Path) -> None:
    bridge(
        folder,
        [
            msg(
                "1700000000555",
                ConversationId=f"{CHANNEL};messageid=1700000000111",
                ConversationType="channel",
            )
        ],
    )
    url = (
        f"{LINK}/19%3Achannel%40thread.tacv2/1700000000555"
        "?tenantId=t&groupId=g&parentMessageId=1700000000111&teamName=MB32829&channelName=General"
    )

    ref = make(folder).parse_reference(url)

    assert ref is not None
    assert ref.external_id == f"{CHANNEL};messageid=1700000000111"
    assert ref.display_ref == "Channel conversation" and ref.locator == url


def test_unencoded_path_segments_work_too(folder: Path) -> None:
    bridge(folder, [msg("1700000000001")])

    ref = make(folder).parse_reference(f"{LINK}/{CHAT}/1700000000001")

    assert ref is not None and ref.external_id == CHAT


def test_the_teams_prefix_form(folder: Path) -> None:
    bridge(folder, [msg("m1", ConversationName="Torque chat")])
    connector = make(folder)

    known = connector.parse_reference(f" teams:{CHAT} ")
    unknown = connector.parse_reference("TEAMS:19:other@thread.v2")

    assert known == SourceRef("teams_thread", CHAT, display_ref="Torque chat", locator=f"teams:{CHAT}")
    assert unknown == SourceRef(
        "teams_thread", "19:other@thread.v2", display_ref=None, locator="TEAMS:19:other@thread.v2"
    )
    assert connector.parse_reference("teams:") is None


@pytest.mark.parametrize(
    "text",
    [
        "",
        "OIL-1",
        "outlook:selected",
        "<abc@acme.com>",
        "https://acme.sharepoint.com/sites/HW/Shared%20Documents/a.docx",
        "https://example.com/l/message/19%3Aabc/1700000000001",
        "http://teams.microsoft.com/l/message/19%3Aabc/1700000000001",
        "ftp://teams.microsoft.com/l/message/19%3Aabc/1700000000001",
        "C:\\Users\\alex\\teams.txt",
    ],
)
def test_other_text_is_not_a_reference(text: str, folder: Path) -> None:
    assert make(folder).parse_reference(text) is None


def test_other_text_needs_no_bridge_folder(tmp_path: Path) -> None:
    assert make(tmp_path / "missing").parse_reference("OIL-1") is None


@pytest.mark.parametrize(
    "url",
    [
        "https://teams.microsoft.com/l/chat/19%3Aabc%40thread.v2/0?tenantId=t",
        "https://teams.microsoft.com/l/meetup-join/19%3Ameeting/0",
        "https://teams.microsoft.com/l/message/19%3Aabc%40thread.v2",
        "https://teams.microsoft.com/",
        "https://teams.microsoft.com/_#/conversations/19:abc?ctx=chat",
    ],
)
def test_another_teams_link_is_rejected_with_a_hint(url: str, folder: Path) -> None:
    with pytest.raises(ValidationError, match="only Teams message links can be linked"):
        make(folder).parse_reference(url)


def test_a_message_link_needs_the_bridge_folder(tmp_path: Path) -> None:
    with pytest.raises(ValidationError, match="Teams bridge folder not found"):
        make(tmp_path / "missing").parse_reference(f"{LINK}/{CHAT_ENCODED}/1")


def test_resolve_describes_the_conversation_without_reading_the_bridge(tmp_path: Path) -> None:
    connector = make(tmp_path / "missing")
    ref = SourceRef("teams_thread", CHAT, display_ref="Torque chat", locator=f"teams:{CHAT}")

    assert connector.resolve(ref) == ObjectMeta(
        "teams_thread", CHAT, title="Torque chat", display_ref="Torque chat", locator=f"teams:{CHAT}"
    )


def test_resolve_rejects_other_object_types(folder: Path) -> None:
    with pytest.raises(ValidationError, match="cannot resolve a file reference"):
        make(folder).resolve(SourceRef("file", "/a.txt"))


@pytest.mark.parametrize("host", ["teams.microsoft.com", "teams.cloud.microsoft", "Teams.Cloud.Microsoft"])
def test_both_teams_hosts_are_handled_alike(host: str, folder: Path) -> None:
    bridge(folder, [msg("1700000000001", ConversationName="Torque chat")])
    connector = make(folder)
    found = f"https://{host}/l/message/{CHAT_ENCODED}/1700000000001?tenantId=t"
    fallback = f"https://{host}/l/message/19%3Aunseen%40thread.v2/1700000000099"

    assert connector.parse_reference(found) == SourceRef(
        "teams_thread", CHAT, display_ref="Torque chat", locator=found
    )
    assert connector.parse_reference(fallback) == SourceRef(
        "teams_thread", "19:unseen@thread.v2", display_ref=None, locator=fallback
    )
    with pytest.raises(ValidationError, match="only Teams message links can be linked"):
        connector.parse_reference(f"https://{host}/l/chat/19%3Aabc%40thread.v2/0")


@pytest.mark.parametrize(
    "host", ["teams.cloud.microsoft.evil.com", "cloud.microsoft", "example.microsoft.com"]
)
def test_other_hosts_are_not_teams_links(host: str, folder: Path) -> None:
    assert make(folder).parse_reference(f"https://{host}/l/message/19%3Aabc/1700000000001") is None


# --- searching and scanning ------------------------------------------------------------------------------


@pytest.fixture
def populated(folder: Path) -> TeamsConnector:
    bridge(
        folder,
        [
            msg(
                "a1",
                ConversationId=CHAT,
                ConversationName="Torque chat",
                Timestamp="2026-03-01T09:00:00+00:00",
            ),
            msg(
                "a2",
                ConversationId=CHAT,
                SenderName="Bob",
                Text="<p>L301 inductor delayed</p>",
                Timestamp="2026-03-02T09:00:00+00:00",
            ),
            msg(
                "b1",
                ConversationId=OTHER_CHAT,
                SenderName="Carla Schmidt",
                Text="Gate review moved",
                Permalink="https://teams.microsoft.com/l/message/x/b1",
                Timestamp="2026-03-05T09:00:00+00:00",
            ),
            msg(
                "c1",
                ConversationId=CHANNEL,
                ConversationType="channel",
                ConversationName="MB32829 General",
                SenderName="Dan",
                SenderEmail="",
                Text="PPAP submitted",
                Timestamp="2026-03-03T09:00:00+00:00",
            ),
        ],
    )
    return make(folder)


def titles(candidates: list[ConversationCandidate]) -> list[str]:
    return [candidate.title for candidate in candidates]


def test_find_matches_name_sender_and_message_text_newest_first(populated: TeamsConnector) -> None:
    assert titles(populated.find_conversations("TORQUE")) == ["Torque chat"]
    assert titles(populated.find_conversations("carla")) == ["Chat with Carla Schmidt"]
    assert titles(populated.find_conversations("l301")) == ["Torque chat"]
    assert titles(populated.find_conversations("ppap")) == ["MB32829 General"]
    assert titles(populated.find_conversations("jane")) == ["Torque chat"]
    assert populated.find_conversations("nothing like this") == []


def test_find_matches_text_after_html_conversion(populated: TeamsConnector) -> None:
    assert populated.find_conversations("<p>") == []
    assert titles(populated.find_conversations("inductor delayed")) == ["Torque chat"]


def test_find_lists_every_conversation_newest_first_up_to_the_limit(populated: TeamsConnector) -> None:
    assert titles(populated.find_conversations("")) == [
        "Chat with Carla Schmidt",
        "MB32829 General",
        "Torque chat",
    ]
    assert titles(populated.find_conversations("", limit=2)) == ["Chat with Carla Schmidt", "MB32829 General"]


def test_a_candidate_carries_a_reference_the_latest_message_and_the_sender(populated: TeamsConnector) -> None:
    candidates = populated.find_conversations("torque")

    assert candidates == [
        ConversationCandidate(
            ref=SourceRef("teams_thread", CHAT, display_ref="Torque chat", locator=f"teams:{CHAT}"),
            title="Torque chat",
            latest=at(2, 9),
            sender="Bob",
            message_count=2,
        )
    ]
    [carla] = populated.find_conversations("carla")
    assert carla.ref.locator == "https://teams.microsoft.com/l/message/x/b1"
    [dan] = populated.find_conversations("ppap")
    assert dan.ref.external_id == CHANNEL and dan.sender == "Dan"


def test_recent_conversations_since_a_moment_newest_first(populated: TeamsConnector) -> None:
    assert titles(populated.recent_conversations(since=at(3, 9))) == [
        "Chat with Carla Schmidt",
        "MB32829 General",
    ]
    assert titles(populated.recent_conversations(since=at(3, 9, 1))) == ["Chat with Carla Schmidt"]
    assert titles(populated.recent_conversations(since=at(1))) == [
        "Chat with Carla Schmidt",
        "MB32829 General",
        "Torque chat",
    ]
    assert populated.recent_conversations(since=at(6)) == []


def test_recent_conversations_takes_a_naive_moment_as_utc(populated: TeamsConnector) -> None:
    assert titles(populated.recent_conversations(since=datetime(2026, 3, 4))) == ["Chat with Carla Schmidt"]


def test_search_candidates_wrap_the_hits_as_object_metadata(populated: TeamsConnector) -> None:
    assert populated.search_candidates(CandidateQuery("torque")) == [
        ObjectMeta(
            "teams_thread", CHAT, title="Torque chat", display_ref="Torque chat", locator=f"teams:{CHAT}"
        )
    ]
    assert [m.external_id for m in populated.search_candidates(CandidateQuery("", limit=2))] == [
        OTHER_CHAT,
        CHANNEL,
    ]


def test_search_candidates_honour_since(populated: TeamsConnector) -> None:
    query = CandidateQuery("", since="2026-03-03T00:00:00.000Z")

    assert [m.external_id for m in populated.search_candidates(query)] == [OTHER_CHAT, CHANNEL]


# --- check -----------------------------------------------------------------------------------------------


def test_check_reports_folder_files_table_and_the_newest_message(folder: Path) -> None:
    bridge(
        folder, [msg("m1"), msg("m2", Timestamp="2026-03-01T11:00:00+00:00")], name="Teams_Bridge_2025.xlsx"
    )
    bridge(folder, [msg("m3", Timestamp="2026-03-02T11:00:00+00:00")])

    report = make(folder).check()

    assert report.ok
    assert [(c.name, c.ok, c.detail) for c in report.checks] == [
        ("folder", True, str(folder)),
        ("files", True, "2 file(s) matching Teams_Bridge*.xlsx"),
        ("file Teams_Bridge_2025.xlsx", True, "table TeamsMessages: 2 message(s)"),
        ("file Teams_Bridge_2026.xlsx", True, "table TeamsMessages: 1 message(s)"),
        ("newest message", True, timeutil.local_display("2026-03-02T11:00:00.000Z")),
    ]


def test_check_reports_problems_per_file_offline_and_conflict_copies(folder: Path, fs) -> None:
    offline = bridge(folder, [msg("m1")], name="Teams_Bridge_2024.xlsx")
    bridge(folder, [msg("m2")], name="Teams_Bridge_2025.xlsx", table="Other")
    bridge(folder, [msg("m3", Timestamp=datetime(2026, 3, 1))], name="Teams_Bridge_2026.xlsx")
    bridge(folder, [msg("m4")], name="Teams_Bridge_2026-PC2.xlsx")
    fs.offline.add(str(offline))

    report = make(folder, fs, hydrate=False).check()

    assert not report.ok
    assert [(c.name, c.ok, c.detail) for c in report.checks] == [
        ("folder", True, str(folder)),
        ("files", True, "3 file(s) matching Teams_Bridge*.xlsx"),
        ("file Teams_Bridge_2024.xlsx", False, OFFLINE_SKIPPED),
        ("file Teams_Bridge_2025.xlsx", False, "table TeamsMessages not found"),
        (
            "file Teams_Bridge_2026.xlsx",
            True,
            "table TeamsMessages: 1 message(s); timestamps without a time zone were taken as UTC",
        ),
        ("conflict copy Teams_Bridge_2026-PC2.xlsx", False, "conflict copy of a bridge file (not read)"),
        ("newest message", True, timeutil.local_display("2026-03-01T00:00:00.000Z")),
    ]


def test_check_notes_a_hydrated_file(folder: Path, fs) -> None:
    path = bridge(folder, [msg("m1")])
    fs.offline.add(str(path))

    [_, _, file_check, _] = make(folder, fs).check().checks

    assert file_check.ok
    assert (
        file_check.detail
        == "table TeamsMessages: 1 message(s); hydrated: downloaded on read (Files On-Demand)"
    )


def test_check_without_a_matching_file(folder: Path) -> None:
    (folder / "other.xlsx").write_bytes(b"x")

    report = make(folder).check()

    assert not report.ok
    assert [(c.name, c.ok, c.detail) for c in report.checks] == [
        ("folder", True, str(folder)),
        ("files", False, f"no file matching Teams_Bridge*.xlsx in {folder}"),
        ("newest message", True, "none captured yet"),
    ]


def test_check_without_the_folder(tmp_path: Path) -> None:
    report = make(tmp_path / "missing").check()

    assert not report.ok
    assert [(c.name, c.ok, c.detail) for c in report.checks] == [
        ("folder", False, f"folder not found: {tmp_path / 'missing'}")
    ]


# --- end to end through the store and the sync engine ----------------------------------------------------


@pytest.fixture
def linked(engine: sqlalchemy.Engine, tmp_path: Path, folder: Path):
    home = tmp_path / "sync_home"
    write_default_config(home)
    bridge(folder, [msg("1700000000001", Text="<p>Lead time for L301?</p>")])
    settings = {"folder": str(folder)}
    with engine.begin() as conn:
        me = people.create_person(conn, "Alex", is_me=True)
        people.add_identity(conn, me.id, "email", "alex@example.com")
        projects.create_project(conn, "OIL", "Open issues")
        issue = issues.create_issue(conn, "OIL", "Inductor delivery", "Chase the supplier", actor_id=me.id)
        row = registry.create_connector(conn, "teams", "Teams", settings=settings)
        connector = TeamsConnector(config(settings))
        ref = connector.parse_reference(f"{LINK}/{CHAT_ENCODED}/1700000000001?tenantId=t")
        assert ref is not None
        object_id = store.upsert_object(conn, row.id, connector.resolve(ref))
        linking.link_source(conn, issue.key, object_id, actor_id=me.id)

    def sync():
        return run_sync(
            engine,
            SyncScope("source", object_id),
            config=load_config(home),
            connector_factory=lambda connector_row: TeamsConnector(config(settings)),
        )

    return engine, object_id, folder, sync


def test_sync_stores_only_new_entries_and_no_versions_and_finds_them_by_text(linked) -> None:
    engine, object_id, folder, sync = linked

    first = sync()
    unchanged = sync()
    bridge(
        folder,
        [
            msg("1700000000001", Text="<p>Lead time for L301?</p>"),
            msg(
                "1700000000002",
                SenderName="Alex",
                SenderEmail="Alex@Example.com",
                Text="<p>Twelve weeks.</p>",
                Timestamp="2026-03-01T11:00:00+00:00",
            ),
        ],
        mtime=MTIME + 60,
    )
    grown = sync()

    assert (first.status, first.results[0].outcome) == ("ok", "new")
    assert (first.stats["new_versions"], first.stats["new_entries"]) == (0, 1)
    assert unchanged.results[0].outcome == "unchanged"
    assert (unchanged.stats["new_versions"], unchanged.stats["new_entries"]) == (0, 0)
    assert (grown.results[0].outcome, grown.stats["new_versions"], grown.stats["new_entries"]) == (
        "changed",
        0,
        1,
    )
    with engine.connect() as conn:
        obj = store.get_object(conn, object_id)
        versions = store.list_versions(conn, object_id)
        entries = store.list_entries(conn, object_id)
        hits = fts.issue_ids_matching(conn, "twelve")
    assert (obj.object_type, obj.external_id, obj.title, obj.fetch_status) == (
        "teams_thread",
        CHAT,
        "Chat with Jane Doe, Alex",
        "ok",
    )
    assert obj.cursor is not None and json.loads(obj.cursor) == {"latest": "2026-03-01T11:00:00.000Z"}
    assert versions == []
    assert [(e.entry_type, e.external_id, e.direction) for e in entries] == [
        ("teams_message", "1700000000001", "in"),
        ("teams_message", "1700000000002", "out"),
    ]
    assert hits


def test_sync_marks_a_removed_conversation_gone_and_keeps_its_content(linked) -> None:
    engine, object_id, folder, sync = linked
    sync()
    os.remove(folder / "Teams_Bridge_2026.xlsx")
    bridge(folder, [msg("z1", ConversationId=OTHER_CHAT)], name="Teams_Bridge_2027.xlsx")

    report = sync()

    assert report.results[0].outcome == "gone"
    with engine.connect() as conn:
        assert store.get_object(conn, object_id).fetch_status == "gone"
        assert len(store.list_entries(conn, object_id)) == 1


def test_sync_reports_a_missing_folder_as_an_error_and_keeps_the_object_intact(linked) -> None:
    engine, object_id, folder, sync = linked
    sync()
    for path in folder.iterdir():
        path.unlink()
    folder.rmdir()

    report = sync()

    assert report.status == "failed" and report.results[0].outcome == "error"
    assert "Teams bridge folder not found" in (report.results[0].error or "")
    with engine.connect() as conn:
        assert store.get_object(conn, object_id).fetch_status == "error"
        assert len(store.list_entries(conn, object_id)) == 1
